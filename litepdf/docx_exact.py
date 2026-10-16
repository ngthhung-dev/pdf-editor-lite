"""PDF -> Word giữ đúng bố cục: mỗi trang Word đúng khổ trang PDF; hình, nét vẽ, nền thành một ảnh nền nằm dưới chữ;
mỗi dòng chữ là một khung chữ đặt đúng toạ độ, đúng phông, cỡ, đậm, nghiêng, màu. Chữ vẫn sửa được trong Word.

Hợp với tờ rơi, hồ sơ năng lực, giấy tờ có bố cục cố định (cách chuyển thường của pdf2docx dàn lại chữ theo dòng chảy
nên hay lệch, mất hình). Đổi lại: hình và nét vẽ không sửa riêng được, văn bản dài không tự chảy sang dòng khác."""
from __future__ import annotations

import io
import re
import unicodedata
from xml.sax.saxutils import escape

import pymupdf

from . import tools

EMU = 12700                 # 1 điểm = 12700 EMU
WORD_MAX = 22 * 72          # Word chỉ cho khổ giấy tối đa 22 inch mỗi chiều

# phông có sẵn trên Windows: giữ đúng tên; phông lạ đổi sang phông gần giống và kéo giãn chữ cho đúng bề rộng
_KNOWN = {"arial": "Arial", "timesnewroman": "Times New Roman", "times": "Times New Roman", "calibri": "Calibri",
          "cambria": "Cambria", "tahoma": "Tahoma", "verdana": "Verdana", "segoeui": "Segoe UI",
          "couriernew": "Courier New", "courier": "Courier New", "georgia": "Georgia", "garamond": "Garamond",
          "helvetica": "Arial", "arialnarrow": "Arial Narrow", "trebuchetms": "Trebuchet MS",
          "comicsansms": "Comic Sans MS", "palatinolinotype": "Palatino Linotype", "bookantiqua": "Book Antiqua",
          "centurygothic": "Century Gothic", "symbol": "Symbol", "wingdings": "Wingdings"}
_METRIC = {"Arial": "he", "Times New Roman": "ti", "Courier New": "co"}      # phông chuẩn PDF cùng số đo
_fonts: dict = {}


def _clean_name(raw: str) -> str:
    name = raw.split("+", 1)[-1]
    name = re.split(r"[,-]", name)[0]
    name = re.sub(r"(PSMT|PS|MT)$", "", name)
    for suf in ("Bold", "Italic", "Oblique", "Regular", "Semibold", "Light", "Medium"):
        name = re.sub(suf + "$", "", name)
    return name


def _word_font(raw: str, flags: int) -> tuple[str, bool]:
    """(tên phông dùng trong Word, có phải phông lạ đã đổi không)."""
    key = re.sub(r"[^a-z]", "", _clean_name(raw).lower())
    if key in _KNOWN:
        return _KNOWN[key], False
    for k, v in _KNOWN.items():
        if key.startswith(k) and len(k) > 4:
            return v, False
    serif = bool(flags & 4) and "sans" not in key
    return ("Times New Roman" if serif else "Arial"), True


def _metric_font(word_font: str, bold: bool, italic: bool):
    base = _METRIC.get(word_font)
    if not base:
        return None
    code = {"he": ("helv", "hebo", "heit", "hebi"), "ti": ("tiro", "tibo", "tiit", "tibi"),
            "co": ("cour", "cobo", "coit", "cobi")}[base][(1 if bold else 0) + (2 if italic else 0)]
    if code not in _fonts:
        _fonts[code] = pymupdf.Font(code)
    return _fonts[code]


def _text_width(font, text: str, size: float) -> float:
    # phông chuẩn PDF không có chữ tiếng Việt có dấu: đo theo chữ gốc không dấu (bề rộng gần như nhau)
    plain = "".join(ch for ch in unicodedata.normalize("NFD", text) if not unicodedata.combining(ch))
    plain = plain.replace("đ", "d").replace("Đ", "D")
    return font.text_length(plain, fontsize=size)


def _hex(color: int) -> str:
    return f"{color & 0xFFFFFF:06X}"


def _run_xml(text: str, font: str, size_hp: int, bold: bool, italic: bool, color: str, scale: int,
             spacing: int = 0) -> str:
    props = [f'<w:rFonts w:ascii="{escape(font)}" w:hAnsi="{escape(font)}" w:eastAsia="{escape(font)}" '
             f'w:cs="{escape(font)}"/>']
    if bold:
        props.append("<w:b/><w:bCs/>")
    if italic:
        props.append("<w:i/><w:iCs/>")
    if color != "000000":
        props.append(f'<w:color w:val="{color}"/>')
    if spacing:
        props.append(f'<w:spacing w:val="{spacing}"/>')
    if scale != 100:
        props.append(f'<w:w w:val="{max(1, min(600, scale))}"/>')
    props.append(f'<w:sz w:val="{size_hp}"/><w:szCs w:val="{size_hp}"/>')
    return f'<w:r><w:rPr>{"".join(props)}</w:rPr><w:t xml:space="preserve">{escape(text)}</w:t></w:r>'


_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
       'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"')


def _anchor(oid: int, x: float, y: float, w: float, h: float, z: int, behind: bool, graphic: str) -> str:
    return (f'<w:r {_NS}><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{z}" behindDoc="{1 if behind else 0}" locked="0" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{int(x * EMU)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{int(y * EMU)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{max(1, int(w * EMU))}" cy="{max(1, int(h * EMU))}"/>'
            f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{oid}" name="Obj{oid}"/><wp:cNvGraphicFramePr/>{graphic}</wp:anchor></w:drawing></w:r>')


def _picture(oid: int, rid: str, w: float, h: float) -> str:
    return ('<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
            f'<pic:nvPicPr><pic:cNvPr id="{oid}" name="Trang{oid}"/><pic:cNvPicPr/></pic:nvPicPr>'
            f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
            f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{int(w * EMU)}" cy="{int(h * EMU)}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>')


def _textbox(runs_xml: str, line_h: float, w: float) -> str:
    twips = max(20, int(round(line_h * 20)))
    return ('<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:xfrm><a:off x="0" y="0"/>'
            f'<a:ext cx="{int(w * EMU)}" cy="{int(line_h * EMU)}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
            '<wps:txbx><w:txbxContent><w:p><w:pPr><w:spacing w:before="0" w:after="0" '
            f'w:line="{twips}" w:lineRule="exact"/><w:ind w:left="0" w:right="0"/></w:pPr>{runs_xml}</w:p>'
            '</w:txbxContent></wps:txbx><wps:bodyPr rot="0" wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" '
            'anchor="t" anchorCtr="0" upright="1"><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic>')


def _lines(page) -> list:
    """Các dòng chữ ngang nhìn thấy được: [(khung, [span...])]."""
    out = []
    d = tools.text_dict(page)
    for b in d.get("blocks", []):
        for ln in b.get("lines", []):
            dx, dy = ln.get("dir", (1, 0))
            if abs(dy) > 0.01 or dx < 0:
                continue
            spans = [s for s in ln.get("spans", []) if s["text"].strip() and s.get("alpha", 255) > 0]
            # khoảng trống lớn giữa hai mẩu chữ (cột, ô của bảng): tách thành khung riêng để mỗi phần đứng đúng chỗ
            groups: list[list] = []
            for sp in spans:
                if groups and pymupdf.Rect(sp["bbox"]).x0 - pymupdf.Rect(groups[-1][-1]["bbox"]).x1 <= 1.2 * sp["size"]:
                    groups[-1].append(sp)
                else:
                    groups.append([sp])
            for g in groups:
                r = pymupdf.Rect(g[0]["bbox"])
                for sp in g[1:]:
                    r |= pymupdf.Rect(sp["bbox"])
                out.append((r, g))
    return out


def _background(doc, pno: int, lines: list, dpi: int) -> tuple[bytes, str]:
    """Ảnh trang không có những dòng chữ sẽ đặt lại thành khung chữ."""
    tmp = pymupdf.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno)
    pg = tmp[0]
    for r, spans in lines:
        for s in spans:
            sb = pymupdf.Rect(s["bbox"])
            o = pymupdf.Point(s["origin"])
            band = pymupdf.Rect(sb.x0 + 0.2, o.y - 0.5 * s["size"], sb.x1 - 0.2, o.y - 0.2 * s["size"])
            pg.add_redact_annot(band if not band.is_empty else sb, fill=False, cross_out=False)
    if lines:
        pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                            text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    rect = pg.rect
    z = dpi / 72
    z = min(z, 9000 / max(rect.width, rect.height))         # không để ảnh quá lớn (trang khổ to)
    pix = pg.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False)
    return pix.tobytes("jpeg", jpg_quality=90), "jpeg"


def convert(doc, out: str, dpi: int = 200, progress=None) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.oxml import parse_xml
    from docx.shared import Pt

    wd = Document()
    for p in list(wd.paragraphs):
        p._element.getparent().remove(p._element)
    oid = 1
    n = doc.page_count
    for pno in range(n):
        page = doc[pno]
        if progress:
            progress(pno, n)
        rot = page.rotation
        lines = _lines(page) if rot == 0 else []       # trang bị xoay: giữ cả trang thành ảnh cho chắc đúng
        W, H = page.rect.width, page.rect.height
        s = min(1.0, WORD_MAX / max(W, H))
        sec = wd.sections[0] if pno == 0 else wd.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Pt(W * s), Pt(H * s)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Pt(0)
        sec.header_distance = sec.footer_distance = Pt(0)
        para = wd.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = Pt(1)
        img, _ext = _background(doc, pno, lines, dpi)
        rid, _image = wd.part.get_or_add_image(io.BytesIO(img))
        z = 1
        para._p.append(parse_xml(_anchor(oid, 0, 0, W * s, H * s, z, True, _picture(oid, rid, W * s, H * s))))
        oid += 1
        for r, spans in lines:
            # từng mẩu chữ: (chữ, kiểu chữ, độ giãn) + khoảng trống phải có trước nó (điểm) để mẩu đó bắt đầu đúng chỗ
            pieces = []
            prev = None
            for sp in spans:
                text = sp["text"]
                sb = pymupdf.Rect(sp["bbox"])
                bold, italic = bool(sp["flags"] & 16), bool(sp["flags"] & 2)
                font, swapped = _word_font(sp["font"], sp["flags"])
                size_hp = max(2, int(round(sp["size"] * s * 2)))
                mf = _metric_font(font, bold, italic)
                scale = 100
                if mf is not None:             # kéo giãn chữ cho đúng bề rộng của mẩu gốc
                    want = sb.width * s
                    have = _text_width(mf, text, size_hp / 2)
                    if have > 1 and want > 1:
                        k = want / have
                        if abs(k - 1) > 0.02 or swapped:
                            scale = int(round(k * 100))
                gap = (sb.x0 - prev[1]) * s if prev is not None else 0.0
                pieces.append({"text": text, "font": font, "hp": size_hp, "b": bold, "i": italic,
                               "c": _hex(sp["color"]), "w": scale, "gap": gap, "mf": mf})
                prev = (sp, sb.x1)
            runs = []
            for k, pc in enumerate(pieces):
                nxt = pieces[k + 1] if k + 1 < len(pieces) else None
                text, extra = pc["text"], 0.0
                if nxt is not None and nxt["gap"] > 0.04 * nxt["hp"] / 2:
                    # khoảng cách tới mẩu sau (chữ dàn đều hai bên, cột…): dồn vào ký tự cuối của mẩu này
                    if not text.endswith(" ") and not nxt["text"].startswith(" "):
                        text += " "
                        sw = _text_width(pc["mf"], " ", pc["hp"] / 2) * pc["w"] / 100 if pc["mf"] else 0.25 * pc["hp"] / 2
                        extra = nxt["gap"] - sw
                    else:
                        extra = nxt["gap"]
                if extra and len(text) > 1:
                    runs.append(_run_xml(text[:-1], pc["font"], pc["hp"], pc["b"], pc["i"], pc["c"], pc["w"]))
                    runs.append(_run_xml(text[-1], pc["font"], pc["hp"], pc["b"], pc["i"], pc["c"], pc["w"],
                                         int(round(extra * 20))))
                else:
                    runs.append(_run_xml(text, pc["font"], pc["hp"], pc["b"], pc["i"], pc["c"], pc["w"],
                                         int(round(extra * 20))))
            z += 1
            top = r.y0 * s
            line_h = max(r.height * s, 1.0)
            # khung rộng hơn dòng một chút: phông trên máy khác đôi chút cũng không bị xuống dòng
            bw = r.width * s * 1.15 + 4
            graphic = _textbox("".join(runs), line_h, bw)
            para._p.append(parse_xml(_anchor(oid, r.x0 * s, top, bw, line_h, z, False, graphic)))
            oid += 1
    wd.save(out)
