"""Lõi xử lý PDF của PDF Editor Lite (không phụ thuộc giao diện)."""
from __future__ import annotations

import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from functools import lru_cache
from pathlib import Path

import pymupdf

pymupdf.TOOLS.mupdf_display_errors(False)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".jfif"}
OFFICE_EXT = {".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".odt", ".ods", ".odp", ".rtf"}
MUPDF_EXT = {".txt", ".html", ".htm", ".epub", ".xps", ".svg", ".fb2", ".cbz", ".mobi", ".md"}
ALL_INPUT_EXT = {".pdf"} | IMAGE_EXT | OFFICE_EXT | MUPDF_EXT


class PdfError(Exception):
    pass


# ----------------------------------------------------------------- khoảng trang
def parse_ranges(spec: str, n: int) -> list[int]:
    """'1-3, 5, 8-' -> [0,1,2,4,7,...]. Số trang tính từ 1."""
    out: list[int] = []
    spec = (spec or "").strip()
    if not spec:
        return list(range(n))
    for part in re.split(r"[,;\s]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d*)\s*-\s*(\d*)", part)
        if m:
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
            step = 1 if b >= a else -1
            rng = range(a, b + step, step)
        elif part.isdigit():
            rng = [int(part)]
        else:
            raise PdfError(f"Không hiểu khoảng trang: “{part}”")
        for p in rng:
            if not 1 <= p <= n:
                raise PdfError(f"Trang {p} nằm ngoài tài liệu (1–{n})")
            out.append(p - 1)
    return out


# ----------------------------------------------------------------------- phông
def _font_dirs() -> list[str]:
    if sys.platform.startswith("win"):
        return [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
                os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts")]
    if sys.platform == "darwin":
        return ["/System/Library/Fonts/Supplemental", "/System/Library/Fonts", "/Library/Fonts",
                os.path.expanduser("~/Library/Fonts")]
    return ["/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
            os.path.expanduser("~/.local/share/fonts")]


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


@lru_cache(maxsize=1)
def font_index() -> dict[str, str]:
    idx: dict[str, str] = {}
    for d in _font_dirs():
        if not d or not os.path.isdir(d):
            continue
        for root, _dirs, files in os.walk(d):
            for f in files:
                stem, ext = os.path.splitext(f)
                if ext.lower() in (".ttf", ".otf"):
                    idx.setdefault(_norm(stem), os.path.join(root, f))
    return idx


_SERIF = ["timesnewroman", "times", "liberationserif", "tinos", "dejavuserif", "notoserif", "freeserif"]
_SANS = ["arial", "liberationsans", "arimo", "dejavusans", "notosans", "freesans", "tahoma", "segoeui",
         "arialunicode", "arialuni", "notosanscjksc", "notosanssc", "notosansjp", "notosanskr", "notosansthai",
         "leelawadeeui", "leelawad", "malgun", "simhei", "simkai", "wqymicrohei", "notosansarabic", "notosansdevanagari"]
_MONO = ["couriernew", "cour", "liberationmono", "cousine", "dejavusansmono", "freemono"]
_ALIASES = {
    "times": _SERIF, "timesnewroman": _SERIF, "cambria": ["cambria", "caladea"] + _SERIF,
    "georgia": ["georgia"] + _SERIF, "garamond": ["garamond"] + _SERIF, "palatino": ["palatino", "pala"] + _SERIF,
    "arial": _SANS, "helvetica": _SANS, "helv": _SANS,
    "calibri": ["calibri", "carlito"] + _SANS, "tahoma": ["tahoma"] + _SANS,
    "verdana": ["verdana"] + _SANS, "segoeui": ["segoeui"] + _SANS,
    "courier": _MONO, "couriernew": _MONO, "consolas": ["consolas", "consola"] + _MONO,
}
_SUFFIX = {
    0: ["", "regular", "r"],
    1: ["bold", "bd", "b"],
    2: ["italic", "i", "oblique", "it"],
    3: ["bolditalic", "bi", "z", "boldoblique", "bditalic"],
}
_BUILTIN = {0: "notos", 1: "notosbo", 2: "notosit", 3: "notosbi"}


@lru_cache(maxsize=64)
def _load_font(path_or_name: str) -> pymupdf.Font:
    if os.path.sep in path_or_name or os.path.exists(path_or_name):
        return pymupdf.Font(fontfile=path_or_name)
    return pymupdf.Font(path_or_name)


def _covers(font: pymupdf.Font, text: str) -> bool:
    return all(ch.isspace() or font.has_glyph(ord(ch)) for ch in text)


def pick_font(pdf_font: str = "", flags: int = 0, text: str = "", bold=None, italic=None) -> pymupdf.Font:
    return pick_font_file(pdf_font, flags, text, bold, italic)[0]


def pick_font_file(pdf_font: str = "", flags: int = 0, text: str = "", bold=None, italic=None):
    """Như pick_font nhưng trả về (Font, đường dẫn file phông hoặc None nếu là phông dựng sẵn)."""
    """Chọn phông hệ thống gần nhất với phông gốc và đủ ký tự (kể cả tiếng Việt)."""
    name = _norm(pdf_font.split("+")[-1])
    is_bold = bool(flags & 16) or "bold" in name or "black" in name or "heavy" in name
    is_ital = bool(flags & 2) or "italic" in name or "oblique" in name
    if bold is not None:
        is_bold = bold
    if italic is not None:
        is_ital = italic
    style = (1 if is_bold else 0) + (2 if is_ital else 0)
    fam = re.sub(r"(bolditalic|boldoblique|bold|italic|oblique|regular|semibold|medium|light|black|heavy)", "", name)
    fam = re.sub(r"(psmt|mt|ps)$", "", fam)
    families: list[str] = []
    if fam:
        families.append(fam)
        for key, lst in _ALIASES.items():
            if fam.startswith(key) or key in fam:
                families += lst
                break
    if flags & 8 or "mono" in name or "courier" in name:
        families += _MONO
    elif flags & 4 or "serif" in name and "sans" not in name:
        families += _SERIF
    families += _SANS
    idx = font_index()
    seen = set()
    for f in families:
        for suf in _SUFFIX[style]:
            key = f + suf
            if key in seen:
                continue
            seen.add(key)
            path = idx.get(key)
            if not path:
                continue
            try:
                font = _load_font(path)
            except Exception:
                continue
            if _covers(font, text):
                return font, path
    try:
        font = _load_font(_BUILTIN[style])
        if _covers(font, text):
            return font, None
    except Exception:
        pass
    return pymupdf.Font("helv"), None


# ------------------------------------------------------------------- mở tài liệu
FROZEN = bool(getattr(sys, "frozen", False))      # đang chạy từ file .exe đã đóng gói (bản cài đặt)


def app_dir() -> str:
    """Thư mục gốc của ứng dụng: nơi đặt file chạy (bản cài đặt) hoặc thư mục mã nguồn (bản portable)."""
    if FROZEN:
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def resource_dir() -> str:
    """Thư mục tài nguyên kèm theo bộ cài (dữ liệu OCR, LibreOffice gọn...)."""
    own = os.path.join(app_dir(), "resources")
    if FROZEN and not os.path.isdir(own):      # macOS (.app) / Linux: tài nguyên nằm trong gói dữ liệu của bản đóng gói
        alt = os.path.join(getattr(sys, "_MEIPASS", ""), "resources")
        if os.path.isdir(alt):
            return alt
    return own


def has_ms_office() -> bool:
    """Máy Windows có Microsoft Word không (app dùng Office để chuyển file Office sang PDF khi không có LibreOffice)."""
    if not sys.platform.startswith("win"):
        return False
    try:
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Word.Application"))
        return True
    except Exception:
        return False


def find_soffice() -> str | None:
    own = os.path.join(resource_dir(), "LibreOffice", "program", "soffice.exe" if sys.platform.startswith("win") else "soffice")
    if os.path.exists(own):
        return own
    for n in ("soffice", "libreoffice"):
        p = shutil.which(n)
        if p:
            return p
    cands = [r"C:\Program Files\LibreOffice\program\soffice.exe",
             r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
             "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for c in cands:
        if os.path.exists(c):
            return c
    return None


_PS_OFFICE = r"""
$ErrorActionPreference='Stop'
$src=$env:LITEPDF_SRC; $dst=$env:LITEPDF_DST; $kind=$env:LITEPDF_KIND
if($kind -eq 'word'){ $a=New-Object -ComObject Word.Application; $a.Visible=$false
  $d=$a.Documents.Open($src,$false,$true); $d.ExportAsFixedFormat($dst,17); $d.Close($false); $a.Quit() }
elseif($kind -eq 'excel'){ $a=New-Object -ComObject Excel.Application; $a.Visible=$false; $a.DisplayAlerts=$false
  $d=$a.Workbooks.Open($src,0,$true); $d.ExportAsFixedFormat(0,$dst); $d.Close($false); $a.Quit() }
else { $a=New-Object -ComObject PowerPoint.Application
  $d=$a.Presentations.Open($src,$true,$false,$false); $d.SaveAs($dst,32); $d.Close(); $a.Quit() }
"""


def _office_kind(ext: str) -> str:
    if ext in (".doc", ".docx", ".rtf", ".odt"):
        return "word"
    if ext in (".xls", ".xlsx", ".ods"):
        return "excel"
    return "ppt"


def office_to_pdf(src: str, out: str) -> str:
    """Chuyển Word/Excel/PowerPoint sang PDF. Trả về tên công cụ đã dùng."""
    src = os.path.abspath(src)
    out = os.path.abspath(out)
    ext = Path(src).suffix.lower()
    so = find_soffice()
    if so:
        with tempfile.TemporaryDirectory() as td:
            prof = Path(td, "profile").as_uri()
            r = subprocess.run([so, f"-env:UserInstallation={prof}", "--headless", "--convert-to", "pdf",
                                "--outdir", td, src], capture_output=True, timeout=300)
            res = Path(td, Path(src).stem + ".pdf")
            if res.exists():
                shutil.move(str(res), out)
                return "LibreOffice"
            err = r.stderr.decode(errors="ignore")[-300:]
        if not sys.platform.startswith("win"):
            raise PdfError("LibreOffice không chuyển được file này. " + err)
    if sys.platform.startswith("win"):
        env = dict(os.environ, LITEPDF_SRC=src, LITEPDF_DST=out, LITEPDF_KIND=_office_kind(ext))
        try:
            subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", _PS_OFFICE],
                           env=env, capture_output=True, timeout=300,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if os.path.exists(out) and os.path.getsize(out) > 0:
                return "Microsoft Office"
        except Exception:
            pass
    if ext in (".docx", ".xlsx", ".pptx"):
        try:
            d = pymupdf.open(src)
            data = d.convert_to_pdf()
            d.close()
            Path(out).write_bytes(data)
            return "bộ dựng sẵn (bố cục cơ bản)"
        except Exception as e:
            raise PdfError(f"Không chuyển được file Office: {e}")
    raise PdfError("Cần cài LibreOffice (miễn phí) hoặc Microsoft Office để chuyển định dạng này sang PDF.")


PAPER_PT = {"A4": (595.28, 841.89), "A3": (841.89, 1190.55), "A5": (419.53, 595.28), "Letter": (612, 792),
            "Legal": (612, 1008)}


def exif_rotation(path: str) -> int:
    """Góc cần xoay theo chiều kim đồng hồ để ảnh chụp (điện thoại, máy ảnh) đứng đúng chiều."""
    try:
        from PySide6.QtGui import QImageIOHandler, QImageReader
        t = QImageReader(path).transformation()
        T = QImageIOHandler.Transformation
        return {T.TransformationRotate90: 90, T.TransformationRotate180: 180, T.TransformationRotate270: 270}.get(t, 0)
    except Exception:
        return 0


def add_image_page(doc, path: str, size: str = "image", orientation: str = "auto", margin_mm: float = 0,
                   rotate: int = 0, enlarge: bool = True, pixmap=None, fit: str = "fit") -> None:
    """Thêm một trang chứa ảnh. size: 'image' (trang bằng đúng ảnh) hoặc tên khổ giấy.
    rotate: xoay ảnh theo chiều kim đồng hồ (0/90/180/270)."""
    pix = pixmap if pixmap is not None else pymupdf.Pixmap(path)
    dpi = pix.xres if pix.xres and pix.xres > 30 else 96
    w, h = pix.width * 72 / dpi, pix.height * 72 / dpi
    rotate = int(rotate) % 360
    if rotate in (90, 270):
        w, h = h, w
    m = max(0.0, margin_mm) * 72 / 25.4
    if size == "image" or size not in PAPER_PT:
        k = min(1.0, 14000 / max(w, h))
        w, h = w * k, h * k
        pw, ph = w + 2 * m, h + 2 * m
        rect = pymupdf.Rect(m, m, m + w, m + h)
    else:
        pw, ph = PAPER_PT[size]
        land = w > h if orientation == "auto" else orientation == "landscape"
        if land:
            pw, ph = ph, pw
        m = min(m, pw / 3, ph / 3)
        bw, bh = pw - 2 * m, ph - 2 * m
        if fit == "stretch":      # kéo giãn kín vùng trang, chấp nhận méo ảnh
            tw, th = bw, bh
        else:
            k = max(bw / w, bh / h) if fit == "fill" else min(bw / w, bh / h)   # fill: lấp kín, phần dư bị cắt
            if not enlarge and fit == "fit":
                k = min(k, 1.0)
            tw, th = w * k, h * k
        rect = pymupdf.Rect((pw - tw) / 2, (ph - th) / 2, (pw + tw) / 2, (ph + th) / 2)
    page = doc.new_page(width=pw, height=ph)
    kw = {"rotate": (360 - rotate) % 360} if rotate else {}   # PyMuPDF xoay ngược chiều kim đồng hồ
    if fit == "stretch" and size in PAPER_PT:
        kw["keep_proportion"] = False
    if pixmap is not None:
        page.insert_image(rect, pixmap=pixmap, **kw)
    else:
        page.insert_image(rect, filename=path, **kw)


def images_to_pdf(paths: list[str], out: str | None = None, a4: bool = False, size: str | None = None,
                  orientation: str = "auto", margin_mm: float = 0, rotations: list[int] | None = None,
                  enlarge: bool = True, progress=None, fit: str = "fit") -> bytes:
    size = size or ("A4" if a4 else "image")
    doc = pymupdf.open()
    for n, p in enumerate(paths):
        rot = rotations[n] if rotations is not None else exif_rotation(p)
        try:
            frames = 1
            if p.lower().endswith((".tif", ".tiff", ".gif")):
                img = pymupdf.open(p)
                frames = len(img)
                if frames > 1:  # TIFF/GIF nhiều khung: mỗi khung một trang
                    for i in range(frames):
                        add_image_page(doc, p, size, orientation, margin_mm, rot, enlarge, pixmap=img[i].get_pixmap(),
                                       fit=fit)
                img.close()
            if frames == 1:
                add_image_page(doc, p, size, orientation, margin_mm, rot, enlarge, fit=fit)
        except Exception as e:
            raise PdfError(f"Không đọc được ảnh {os.path.basename(p)}: {e}")
        if progress:
            progress(n + 1, len(paths))
    data = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    if out:
        Path(out).write_bytes(data)
    return data


def any_to_pdf_bytes(path: str) -> tuple[bytes, str]:
    """Đọc file bất kỳ được hỗ trợ và trả về (bytes PDF, ghi chú công cụ)."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return Path(path).read_bytes(), ""
    if ext in IMAGE_EXT:
        return images_to_pdf([path]), "ảnh"
    if ext in OFFICE_EXT:
        with tempfile.TemporaryDirectory() as td:
            out = os.path.join(td, "o.pdf")
            eng = office_to_pdf(path, out)
            return Path(out).read_bytes(), eng
    try:
        d = pymupdf.open(path)
        data = d.convert_to_pdf()
        d.close()
        return data, "MuPDF"
    except Exception as e:
        raise PdfError(f"Định dạng không hỗ trợ hoặc file hỏng: {e}")


def open_pdf_bytes(data: bytes, password: str | None = None) -> tuple[pymupdf.Document, bool]:
    """Mở PDF từ bytes; nếu có mật khẩu thì giải mã trong bộ nhớ. Trả về (doc, đã_mã_hoá)."""
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        raise PdfError(f"Không mở được PDF: {e}")
    enc = bool(doc.is_encrypted or doc.needs_pass)
    if doc.needs_pass:
        if password is None or not doc.authenticate(password):
            doc.close()
            raise PermissionError("Cần mật khẩu")
    if enc:
        plain = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_NONE)
        doc.close()
        doc = pymupdf.open(stream=plain, filetype="pdf")
    return doc, enc


# ----------------------------------------------------------------- gộp / tách
def merge_files(paths: list[str], out: str, progress=None) -> int:
    dst = pymupdf.open()
    for i, p in enumerate(paths):
        data, _ = any_to_pdf_bytes(p)
        try:
            src, _enc = open_pdf_bytes(data)
        except PermissionError:
            raise PdfError(f"{os.path.basename(p)} có mật khẩu, hãy mở và gỡ mật khẩu trước.")
        dst.insert_pdf(src)
        src.close()
        if progress:
            progress(i + 1, len(paths))
    n = dst.page_count
    dst.save(out, garbage=3, deflate=True)
    dst.close()
    return n


def _safe_name(text: str, limit: int = 60) -> str:
    text = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", text).strip(" .")
    return re.sub(r"\s+", " ", text)[:limit].strip() or "khong_ten"


def _pages_label(pages: list[int]) -> str:
    """[0,1,2,6] -> '1-3+7'."""
    runs: list[list[int]] = []
    for p in pages:
        if runs and runs[-1][1] + 1 == p:
            runs[-1][1] = p
        else:
            runs.append([p, p])
    return "+".join(f"{a + 1}" if a == b else f"{a + 1}-{b + 1}" for a, b in runs)


def split_plan(doc, mode: str, value="") -> list[tuple[str, list[int]]]:
    """Lập danh sách file sẽ tạo khi tách: [(nhãn, [chỉ số trang...]), ...].
    mode: 'each'  - mỗi trang một file
          'every' - cứ N trang một file (value = N)
          'ranges'- mỗi khoảng một file; các khoảng cách nhau bằng dấu phẩy, chấm phẩy hoặc xuống dòng,
                    ví dụ '1-3, 4-10, 11-'. Muốn ghép các đoạn rời vào chung một file thì nối bằng dấu +: '1-3+7'
          'toc'   - tách theo bookmark (value = cấp bookmark, mặc định 1), mỗi mục một file
          'at'    - cắt tại các trang cho trước: '5, 12' -> 1-4, 5-11, 12-hết"""
    n = doc.page_count
    plan: list[tuple[str, list[int]]] = []
    if mode == "each":
        plan = [(f"trang_{i + 1}", [i]) for i in range(n)]
    elif mode == "every":
        k = int(value or 0)
        if k < 1:
            raise PdfError("Số trang mỗi file phải từ 1 trở lên.")
        for i in range(0, n, k):
            g = list(range(i, min(i + k, n)))
            plan.append(("trang_" + _pages_label(g), g))
    elif mode == "ranges":
        for part in re.split(r"[,;\n]+", str(value)):
            part = part.strip()
            if not part:
                continue
            g: list[int] = []
            for sub in part.split("+"):
                if not sub.strip():
                    raise PdfError(f"Không hiểu khoảng trang: “{part}”")
                g += parse_ranges(sub.strip(), n)
            plan.append(("trang_" + _pages_label(g), g))
        if not plan:
            raise PdfError("Hãy nhập khoảng trang, ví dụ: 1-3, 4-10, 11-")
    elif mode == "at":
        cuts = sorted({p for p in parse_ranges(str(value), n) if p > 0}) if str(value).strip() else []
        if not cuts:
            raise PdfError("Hãy nhập các trang bắt đầu file mới, ví dụ: 5, 12")
        starts = [0] + cuts
        for a, b in zip(starts, starts[1:] + [n]):
            g = list(range(a, b))
            plan.append(("trang_" + _pages_label(g), g))
    elif mode == "toc":
        level = int(value or 1)
        marks: list[tuple[int, str]] = []
        for lvl, title, page in doc.get_toc(simple=True):
            if lvl <= level and 1 <= page <= n and (not marks or marks[-1][0] != page - 1):
                marks.append((page - 1, title))
        if not marks:
            raise PdfError("Tài liệu không có bookmark để tách theo.")
        if marks[0][0] > 0:
            marks.insert(0, (0, "Phần đầu"))
        for k, (start, title) in enumerate(marks):
            end = marks[k + 1][0] if k + 1 < len(marks) else n
            plan.append((_safe_name(title), list(range(start, end))))
    else:
        raise PdfError("Cách tách không hợp lệ.")
    return plan


def split_filenames(plan, base: str) -> list[str]:
    width = max(2, len(str(len(plan))))
    return [f"{base}_{i:0{width}d}_{label}.pdf" for i, (label, _pages) in enumerate(plan, 1)]


def split_write(doc, plan, out_dir: str, base: str, progress=None) -> list[str]:
    """Ghi các file theo kế hoạch tách. Trả về danh sách đường dẫn."""
    outs = []
    names = split_filenames(plan, base)
    for k, ((_label, pages), name) in enumerate(zip(plan, names)):
        d = pymupdf.open()
        runs: list[list[int]] = []
        for p in pages:  # gom các trang liên tiếp để giữ liên kết
            if runs and runs[-1][1] + 1 == p:
                runs[-1][1] = p
            else:
                runs.append([p, p])
        for a, b in runs:
            d.insert_pdf(doc, from_page=a, to_page=b)
        path = os.path.join(out_dir, name)
        d.save(path, garbage=3, deflate=True)
        d.close()
        outs.append(path)
        if progress:
            progress(k + 1, len(plan))
    return outs


def split_doc(doc, out_dir: str, base: str, mode: str, value="") -> list[str]:
    return split_write(doc, split_plan(doc, mode, value), out_dir, base)


# ------------------------------------------------------------- xuất từ PDF
def pdf_to_images(doc, out_dir, base, fmt="png", dpi=150, pages=None, progress=None) -> list[str]:
    pages = pages if pages is not None else list(range(doc.page_count))
    outs = []
    for k, i in enumerate(pages):
        pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
        path = os.path.join(out_dir, f"{base}_trang_{i + 1:03d}.{fmt}")
        if fmt in ("jpg", "jpeg"):
            pix.save(path, jpg_quality=90)
        else:
            pix.save(path)
        outs.append(path)
        if progress:
            progress(k + 1, len(pages))
    return outs


def pdf_to_text(doc, out: str) -> None:
    with open(out, "w", encoding="utf-8") as f:
        for i, page in enumerate(doc):
            f.write(page_text(page, sort=True))
            f.write(f"\n\n----- Trang {i + 1} -----\n\n")


def pdf_to_html(doc, out: str) -> None:
    with open(out, "w", encoding="utf-8") as f:
        f.write("<!doctype html><html><head><meta charset='utf-8'><title>%s</title></head><body>\n"
                % os.path.basename(out))
        for page in _text_ready(doc):
            f.write(page.get_text("html"))
        f.write("</body></html>")


def pdf_to_docx(pdf_path: str, out: str) -> None:
    try:
        from pdf2docx import Converter
    except Exception as e:
        raise PdfError("Thiếu thư viện pdf2docx (chạy lại file cài đặt để bổ sung). " + str(e))
    import logging
    logging.disable(logging.CRITICAL)
    try:
        d = pymupdf.open(pdf_path)
        if d.needs_pass == 0 and add_missing_tounicode(d):
            d.saveIncr()
        d.close()
    except Exception:
        pass
    try:
        cv = Converter(pdf_path)
        cv.convert(out)
        cv.close()
        from . import ooxml_fix
        ooxml_fix.fix(out)
    except Exception as e:
        raise PdfError(f"Không chuyển được sang Word: {e}")
    finally:
        logging.disable(logging.NOTSET)


def pdf_to_xlsx(doc, out: str) -> int:
    """Xuất bảng trong PDF ra Excel. Trả về số bảng tìm thấy."""
    try:
        import openpyxl
    except Exception as e:
        raise PdfError("Thiếu thư viện openpyxl. " + str(e))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    total = 0
    for i, page in enumerate(_text_ready(doc)):
        ws = wb.create_sheet(f"Trang {i + 1}")
        row = 1
        try:
            tabs = page.find_tables().tables
        except Exception:
            tabs = []
        for t in tabs:
            total += 1
            for r in t.extract():
                for c, v in enumerate(r, 1):
                    ws.cell(row=row, column=c, value=_num(v))
                row += 1
            row += 1
        if not tabs:  # không có bảng: ghi theo dòng chữ, tách cột theo khoảng trống lớn
            for line in _lines_as_rows(page):
                for c, v in enumerate(line, 1):
                    ws.cell(row=row, column=c, value=_num(v))
                row += 1
    wb.save(out)
    from . import ooxml_fix
    ooxml_fix.fix(out)
    return total


def _num(v):
    if v is None:
        return None
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s) and len(s) < 15 and not (len(s) > 1 and s.startswith("0")):
        return int(s)
    return s


def _lines_as_rows(page) -> list[list[str]]:
    words = page_words(page, sort=True)
    rows: list[list] = []
    for w in words:
        if rows and abs(rows[-1][0] - w[3]) < 3:
            rows[-1][1].append(w)
        else:
            rows.append([w[3], [w]])
    out = []
    for _y, ws in rows:
        ws.sort(key=lambda w: w[0])
        cells, cur, last = [], [], None
        for w in ws:
            if last is not None and w[0] - last > 18:
                cells.append(" ".join(cur))
                cur = []
            cur.append(w[4])
            last = w[2]
        cells.append(" ".join(cur))
        out.append(cells)
    return out


# ------------------------------------------------------------------- tiện ích
def compress_bytes(doc, level: str = "medium") -> bytes:
    """level: low (giữ chất lượng) | medium | high (file nhỏ nhất)."""
    tmp = pymupdf.open(stream=doc.tobytes(), filetype="pdf")
    cfg = {"low": None, "medium": (200, 150, 75), "high": (120, 96, 55)}[level]
    if cfg:
        try:
            tmp.rewrite_images(dpi_threshold=cfg[0], dpi_target=cfg[1], quality=cfg[2])
        except Exception:
            pass
    try:
        tmp.subset_fonts()
    except Exception:
        pass
    data = tmp.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=False,
                       use_objstms=1)
    tmp.close()
    return data


def save_encrypted(doc, out: str, user_pw: str, owner_pw: str = "", allow_print=True, allow_copy=True,
                   allow_edit=False) -> None:
    perm = pymupdf.PDF_PERM_ACCESSIBILITY
    if allow_print:
        perm |= pymupdf.PDF_PERM_PRINT | pymupdf.PDF_PERM_PRINT_HQ
    if allow_copy:
        perm |= pymupdf.PDF_PERM_COPY
    if allow_edit:
        perm |= pymupdf.PDF_PERM_MODIFY | pymupdf.PDF_PERM_ANNOTATE | pymupdf.PDF_PERM_FORM | pymupdf.PDF_PERM_ASSEMBLE
    doc.save(out, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_pw or None,
             owner_pw=owner_pw or user_pw or None, permissions=perm, garbage=3, deflate=True)


def add_watermark(doc, text: str, pages=None, size=54, color=(0.8, 0.1, 0.1), opacity=0.18, angle=45) -> None:
    font = pick_font("arial", 16, text)
    for i in (pages if pages is not None else range(doc.page_count)):
        page = doc[i]
        r = page.rect  # kích thước đang hiển thị (đã tính xoay trang)
        tl = font.text_length(text, size)
        s = min(size, size * (0.8 * math.hypot(r.width, r.height)) / tl) if tl else size
        tl = font.text_length(text, s)
        c = pymupdf.Point(r.width / 2, r.height / 2) * page.derotation_matrix
        tw = pymupdf.TextWriter(page.rect, color=color)
        tw.append((c.x - tl / 2, c.y + s * 0.35), text, font=font, fontsize=s)
        tw.write_text(page, opacity=opacity, morph=(c, pymupdf.Matrix(angle + page.rotation)), overlay=True)


def add_page_numbers(doc, fmt="{n} / {total}", pos="bottom-center", size=10, start=1, pages=None,
                     color=(0.2, 0.2, 0.2), margin=28) -> None:
    font = pick_font("arial", 0, fmt)
    total = doc.page_count
    for i in (pages if pages is not None else range(total)):
        page = doc[i]
        r = page.rect
        text = fmt.format(n=i + start, total=total + start - 1)
        tl = font.text_length(text, size)
        v, h = pos.split("-")
        y = r.y1 - margin if v == "bottom" else r.y0 + margin + size * 0.7
        x = {"left": r.x0 + margin, "center": r.x0 + (r.width - tl) / 2, "right": r.x1 - margin - tl}[h]
        write_display_text(page, (x, y), text, font, size, color)


def find_tessdata() -> str | None:
    """Tìm thư mục dữ liệu ngôn ngữ của Tesseract: ưu tiên bản kèm theo bộ cài, rồi tự dò các chỗ cài mặc định."""
    own = os.path.join(resource_dir(), "tessdata")
    if os.path.isdir(own) and any(f.endswith(".traineddata") for f in os.listdir(own)):
        os.environ["TESSDATA_PREFIX"] = own
        return own
    env = os.environ.get("TESSDATA_PREFIX")
    if env and os.path.isdir(env):
        return env
    cands = []
    if sys.platform.startswith("win"):
        for base in (os.environ.get("ProgramFiles", r"C:\Program Files"), os.environ.get("ProgramFiles(x86)", ""),
                     os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs"), r"C:\Program Files"):
            if base:
                cands.append(os.path.join(base, "Tesseract-OCR", "tessdata"))
    elif sys.platform == "darwin":
        cands += ["/opt/homebrew/share/tessdata", "/usr/local/share/tessdata"]
    else:
        cands += ["/usr/share/tesseract-ocr/5/tessdata", "/usr/share/tesseract-ocr/4.00/tessdata", "/usr/share/tessdata"]
    for c in cands:
        if os.path.isdir(c) and any(f.endswith(".traineddata") for f in os.listdir(c)):
            os.environ["TESSDATA_PREFIX"] = c
            return c
    try:
        td = pymupdf.get_tessdata()
        if td and os.path.isdir(td):
            return td
    except Exception:
        pass
    return None


def tess_languages() -> list[str]:
    td = find_tessdata()
    if not td:
        return []
    return sorted(f[:-12] for f in os.listdir(td) if f.endswith(".traineddata") and f != "osd.traineddata")


def ocr_bytes(doc, language="eng", dpi=300, progress=None) -> bytes:
    """Tạo PDF tìm kiếm được từ PDF scan (cần cài Tesseract)."""
    if not tess_languages():
        raise PdfError("Chưa cài Tesseract OCR. Cài Tesseract (kèm gói tiếng Việt “vie”) rồi thử lại.")
    out = pymupdf.open()
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        try:
            one = pymupdf.open("pdf", pix.pdfocr_tobytes(language=language, tessdata=find_tessdata()))
        except Exception as e:
            raise PdfError(f"OCR lỗi ở trang {i + 1}: {e}")
        out.insert_pdf(one)
        one.close()
        if progress:
            progress(i + 1, doc.page_count)
    data = out.tobytes(garbage=3, deflate=True)
    out.close()
    return data


# ------------------------------------------------------------- ghi chú dán (sticky note)
NOTE_MARK = "lp|note"
_NOTE_AP = b"""%lpnote
q
1 0.83 0.22 rg 0.60 0.42 0.02 RG 0.8 w 1 j
3 4.5 m 3 1 l 7.2 4.5 l 13 4.5 l 14.1 4.5 15 5.4 15 6.5 c 15 13 l 15 14.1 14.1 15 13 15 c 3 15 l
1.9 15 1 14.1 1 13 c 1 6.5 l 1 5.4 1.9 4.5 3 4.5 c h B
0.42 0.29 0.02 RG 1 w 1 J
4 12.2 m 12 12.2 l S 4 9.8 m 12 9.8 l S 4 7.4 m 9 7.4 l S
Q
"""


def style_notes(page) -> int:
    """Thay hình mặc định (ô vuông thô) của các ghi chú do app tạo bằng biểu tượng bong bóng ghi chú. Trả về số ghi chú đã đổi."""
    doc = page.parent
    n = 0
    try:
        annots = list(page.annots(types=[pymupdf.PDF_ANNOT_TEXT]))
    except Exception:
        return 0
    for a in annots:
        try:
            mine = a.info.get("subject") == NOTE_MARK
            if not mine and not a.info.get("subject"):      # ghi chú tạo bằng bản cũ của app: biểu tượng Comment màu vàng 1/0.82/0
                st = (a.colors or {}).get("stroke") or ()
                mine = doc.xref_get_key(a.xref, "Name")[1] == "/Comment" and len(st) == 3 \
                    and abs(st[0] - 1) < 0.02 and abs(st[1] - 0.82) < 0.02 and st[2] < 0.02
            if not mine:
                continue
            kind, val = doc.xref_get_key(a.xref, "AP/N")
            if kind != "xref":
                continue
            x = int(val.split()[0])
            if doc.xref_stream(x).startswith(b"%lpnote"):
                continue
            doc.xref_set_key(x, "BBox", "[0 0 16 16]")
            doc.update_stream(x, _NOTE_AP)
            n += 1
        except Exception:
            pass
    return n


def page_notes(page) -> list[tuple]:
    """[(khung, loại, nội dung, tác giả, ngày)] của các ghi chú có nội dung trên trang (để hiện khi rê chuột)."""
    out = []
    try:
        for a in page.annots():
            info = a.info
            text = (info.get("content") or "").strip()
            if text and not (info.get("subject") or "").startswith("lp|image"):
                out.append((pymupdf.Rect(a.rect), a.type[1], text, info.get("title") or "", pdf_date(info.get("modDate") or info.get("creationDate"))))
    except Exception:
        pass
    return out


# ------------------------------------------------------------- sửa chữ gốc
def _rgb(c: int) -> tuple[float, float, float]:
    return ((c >> 16) & 255) / 255, ((c >> 8) & 255) / 255, (c & 255) / 255


# ------------------------------------------------- chữ của phông thiếu bảng mã (ToUnicode)
# Nhiều file PDF tiếng Việt (xuất từ phần mềm cũ, cổng dịch vụ công…) nhúng phông TrueType kiểu Identity-H mà không kèm
# bảng mã chữ: trên trang chữ vẫn hiện đúng, nhưng chữ đọc ra là số thứ tự hình chữ ("NLQK" thay cho "kinh"). Bảng mã
# riêng (cmap) của phông nhúng thì vẫn còn: tra ngược số hình chữ -> chữ thật.
def _glyph_map(doc, xref: int) -> dict | None:
    cache = doc.__dict__.setdefault("_lp_glyph_maps", {})
    if xref in cache:
        return cache[xref]
    m = None
    try:
        import io
        from fontTools.agl import toUnicode
        from fontTools.ttLib import TTFont
        buf = doc.extract_font(xref)[3]
        if buf:
            f = TTFont(io.BytesIO(buf), lazy=True)
            order = f.getGlyphOrder()
            m = {}
            for u, g in (f.getBestCmap() or {}).items():
                try:
                    gid = f.getGlyphID(g)
                except Exception:
                    continue
                if gid not in m or chr(u).isprintable() and not chr(m[gid]).isprintable():
                    m[gid] = u
            for gid, name in enumerate(order):     # hình chữ không có trong bảng mã: đoán theo tên (uni1EA1, aacute…)
                if gid not in m and name and not name.startswith((".", "glyph")):
                    t = toUnicode(name)
                    if len(t) == 1:
                        m[gid] = ord(t)
            m = {k: chr(v) for k, v in m.items()} or None
    except Exception:
        m = None
    cache[xref] = m
    return m


def _tounicode_cmap(m: dict) -> bytes:
    items = sorted((g, u) for g, u in m.items() if 0 <= g <= 0xFFFF)
    out = [b"/CIDInit /ProcSet findresource begin 12 dict begin begincmap",
           b"/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def",
           b"/CMapName /Adobe-Identity-UCS def /CMapType 2 def",
           b"1 begincodespacerange <0000> <FFFF> endcodespacerange"]
    for k in range(0, len(items), 100):
        chunk = items[k:k + 100]
        out.append(b"%d beginbfchar" % len(chunk))
        for g, u in chunk:
            out.append(b"<%04X> <%s>" % (g, u.encode("utf-16-be").hex().upper().encode()))
        out.append(b"endbfchar")
    out.append(b"endcmap CMapName currentdict /CMap defineresource pop end end")
    return b"\n".join(out)


def add_missing_tounicode(doc) -> int:
    """Gắn bảng mã ToUnicode (dựng từ bảng mã của phông nhúng) cho các phông thiếu, để công cụ khác
    (pdf2docx, xuất HTML, bảng Excel) đọc đúng chữ. Sửa trên doc truyền vào; trả về số phông đã gắn."""
    n, seen = 0, set()
    for page in doc:
        try:
            fonts = page.get_fonts()
        except Exception:
            continue
        for xref, _ext, typ, _base, _name, enc in fonts:
            if xref in seen or typ != "Type0" or not str(enc).startswith("Identity"):
                continue
            seen.add(xref)
            try:
                if doc.xref_get_key(xref, "ToUnicode")[0] != "null":
                    continue
                desc = doc.xref_get_key(xref, "DescendantFonts")
                dx = int(desc[1].strip("[] ").split()[0]) if desc[0] == "array" else \
                    int(desc[1].split()[0]) if desc[0] == "xref" else 0
                if dx:
                    c2g = doc.xref_get_key(dx, "CIDToGIDMap")
                    if c2g[0] not in ("null", "name"):
                        continue          # CID khác số hình chữ: không đoán
                m = _glyph_map(doc, xref)
                if not m:
                    continue
                sx = doc.get_new_xref()
                doc.update_object(sx, "<<>>")
                doc.update_stream(sx, _tounicode_cmap(m))
                doc.xref_set_key(xref, "ToUnicode", f"{sx} 0 R")
                n += 1
            except Exception:
                continue
    return n


def _text_ready(doc):
    """Bản sao trong bộ nhớ đã gắn ToUnicode nếu cần (không ghi gì ra đĩa); không cần thì trả lại doc gốc."""
    try:
        cp = pymupdf.open("pdf", doc.tobytes())
        if add_missing_tounicode(cp):
            return cp
        cp.close()
    except Exception:
        pass
    return doc


def _broken_fonts(page) -> dict:
    """{tên phông như trong span: bảng số hình chữ -> chữ} cho các phông thiếu bảng mã của trang."""
    doc = page.parent
    out: dict = {}
    try:
        fonts = page.get_fonts()
    except Exception:
        return out
    for xref, _ext, typ, base, _name, enc in fonts:
        if typ != "Type0" or not str(enc).startswith("Identity"):
            continue
        try:
            if doc.xref_get_key(xref, "ToUnicode")[0] != "null":
                continue
        except Exception:
            continue
        m = _glyph_map(doc, xref)
        if m:
            out.setdefault(base.split("+", 1)[-1], m)
    return out


def _fix_spans(page, d: dict) -> dict:
    """Sửa chữ đọc sai của các phông thiếu bảng mã ngay trong kết quả get_text("dict"/"rawdict")."""
    maps = _broken_fonts(page)
    if not maps:
        return d
    spans = [s for b in d.get("blocks", []) for ln in b.get("lines", []) for s in ln.get("spans", [])]
    # chỉ sửa phông thật sự đọc ra chữ sai: có ký tự điều khiển (khoảng trắng bị đọc thành \x03…)
    bad = set()
    for s in spans:
        txt = s.get("text") if "text" in s else "".join(c["c"] for c in s.get("chars", []))
        if s["font"] in maps and any(ord(c) < 32 and c not in "\t\n\r" for c in txt):
            bad.add(s["font"])
    import unicodedata as _ud
    for b in d.get("blocks", []):
        for ln in b.get("lines", []):
            # Số hình chữ trùng với mã của một dấu kết hợp (U+0300…): thư viện đọc nó như dấu (rộng 0) rồi chèn thêm một
            # dấu cách giả đúng bằng bề rộng của chữ ("lầ n"). Gặp vậy thì bỏ dấu cách giả, trả bề rộng lại cho chữ.
            hang = None          # (span, ký tự) vừa đọc ra từ một mã trùng dấu kết hợp
            for s in ln.get("spans", []):
                m = maps.get(s["font"]) if s["font"] in bad else None
                if "chars" in s:
                    keep = []
                    for c in s["chars"]:
                        code = c["c"]
                        if hang is not None and code == " ":
                            hang["bbox"] = tuple(pymupdf.Rect(hang["bbox"]) | pymupdf.Rect(c["bbox"]))
                            hang = None
                            continue
                        hang = None
                        if m:
                            c["c"] = m.get(ord(code), code)
                            if c["c"] != code and _ud.category(code) in ("Mn", "Me"):
                                hang = c
                        keep.append(c)
                    s["chars"] = keep
                if "text" in s:
                    out = []
                    for code in s["text"]:
                        if hang is not None and code == " ":
                            if hang is not s and not out:      # dấu cách giả đứng riêng một mẩu: trả bề rộng cho mẩu trước
                                hang["bbox"] = tuple(pymupdf.Rect(hang["bbox"]) | pymupdf.Rect(s["bbox"]))
                            hang = None
                            continue
                        hang = None
                        ch = m.get(ord(code), code) if m else code
                        if m and ch != code and _ud.category(code) in ("Mn", "Me"):
                            hang = s
                        out.append(ch)
                    s["text"] = "".join(out)
    return d


def text_dict(page) -> dict:
    return _fix_spans(page, page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES))


def page_words(page, sort: bool = False) -> list:
    """Như page.get_text("words") nhưng đọc đúng chữ của phông thiếu bảng mã."""
    if not _broken_fonts(page):
        return page.get_text("words", sort=sort)
    d = _fix_spans(page, page.get_text("rawdict", flags=pymupdf.TEXT_USE_GID_FOR_UNKNOWN_UNICODE))
    words = []
    for bi, b in enumerate(d.get("blocks", [])):
        for li, ln in enumerate(b.get("lines", [])):
            cur, box, wn = "", None, 0
            for s in ln.get("spans", []):
                for c in s.get("chars", []):
                    if c["c"].isspace():
                        if cur:
                            words.append((box.x0, box.y0, box.x1, box.y1, cur, bi, li, wn))
                            wn += 1
                        cur, box = "", None
                        continue
                    r = pymupdf.Rect(c["bbox"])
                    cur += c["c"]
                    box = r if box is None else box | r
            if cur:
                words.append((box.x0, box.y0, box.x1, box.y1, cur, bi, li, wn))
    if sort:
        words.sort(key=lambda w: (round(w[3], 1), w[0]))
    return words


def page_text(page, sort: bool = False) -> str:
    """Như page.get_text("text") nhưng đọc đúng chữ của phông thiếu bảng mã."""
    if not _broken_fonts(page):
        return page.get_text("text", sort=sort)
    d = text_dict(page)
    out = []
    for b in d.get("blocks", []):
        for ln in b.get("lines", []):
            out.append("".join(s["text"] for s in ln.get("spans", [])))
        out.append("")
    return "\n".join(out)


def text_runs(d: dict) -> list[dict]:
    """Các "dòng chữ" sửa được của trang: những mẩu chữ cùng kiểu, cùng hàng, nằm sát nhau được nối thành một dòng.
    (Nhiều file PDF ghi một dòng thành nhiều mẩu rời, có khi từng chữ cái; không nối thì chỉ sửa được từng mẩu.)
    Kết quả được nhớ ngay trong d để rê chuột không phải tính lại."""
    if "_runs" in d:
        return d["_runs"]
    pieces = []
    for bi, block in enumerate(d.get("blocks", [])):
        for line in block.get("lines", []):
            dr = tuple(line.get("dir", (1, 0)))
            for s in line.get("spans", []):
                if s["text"].strip():
                    s["_blk"] = bi
                    pieces.append((s, dr))
    horiz = [(s, dr) for s, dr in pieces if abs(dr[1]) < 0.01 and dr[0] > 0]
    other = [(s, dr) for s, dr in pieces if not (abs(dr[1]) < 0.01 and dr[0] > 0)]
    horiz.sort(key=lambda x: (round(x[0]["origin"][1], 1), x[0]["bbox"][0]))
    runs: list[dict] = []

    def style_of(s):
        # cùng một phông có thể bị file chia thành nhiều phông con (tên có tiền tố "ABCDEF+", hoặc nhiều phông Type3)
        name = s["font"]
        name = "Type3" if name.startswith("Type3") else name.split("+")[-1]
        return (name, s["color"], s["flags"] & 18)

    def start(s, dr):
        runs.append({"bbox": pymupdf.Rect(s["bbox"]), "text": s["text"], "origin": pymupdf.Point(s["origin"]),
                     "size": s["size"], "color": _rgb(s["color"]), "font": s["font"], "flags": s["flags"], "dir": dr,
                     "block": s.get("_blk"),
                     "_parts": {style_of(s): [len(s["text"]), s]},
                     "segs": [[s["text"], bool(s["flags"] & 16), bool(s["flags"] & 2), _rgb(s["color"])]]})
    for s, dr in horiz:
        r = runs[-1] if runs else None
        # Nối theo DÒNG nhìn thấy trên trang: cùng hàng, cỡ chữ gần bằng nhau, đứng sát nhau. Chữ đậm / nghiêng / khác màu
        # nằm giữa dòng vẫn thuộc dòng đó (kiểu chữ của dòng lấy theo phần chiếm nhiều chữ nhất).
        if r is not None and abs(r["origin"].y - s["origin"][1]) < 0.25 * s["size"] \
                and 0.88 <= s["size"] / max(0.1, r["size"]) <= 1.14:
            gap = s["bbox"][0] - r["bbox"].x1
            if -0.6 * s["size"] <= gap <= 1.2 * s["size"]:
                glue = "" if r["text"].endswith(" ") or s["text"].startswith(" ") or gap < 0.18 * s["size"] else " "
                r["text"] += glue + s["text"]
                r["bbox"] |= pymupdf.Rect(s["bbox"])
                seg = [glue + s["text"], bool(s["flags"] & 16), bool(s["flags"] & 2), _rgb(s["color"])]
                if r["segs"][-1][1:] == seg[1:]:
                    r["segs"][-1][0] += seg[0]            # cùng kiểu với mẩu trước: gộp lại
                else:
                    r["segs"].append(seg)
                part = r["_parts"].setdefault(style_of(s), [0, s])
                part[0] += len(s["text"])
                continue
        start(s, dr)
    for r in runs:
        if len(r["_parts"]) > 1:
            main = max(r["_parts"].values(), key=lambda x: x[0])[1]
            r["font"], r["flags"], r["color"], r["size"] = main["font"], main["flags"], _rgb(main["color"]), main["size"]
            r["mixed"] = True                # dòng có nhiều kiểu chữ: viết lại sẽ về một kiểu
    for s, dr in other:              # chữ xoay / chữ dọc: giữ từng mẩu như cũ
        start(s, dr)
    for r in runs:
        r["text"] = r["text"].rstrip()
        r["segs"][-1][0] = r["segs"][-1][0].rstrip()
        r["segs"] = [tuple(x) for x in r["segs"] if x[0]]
    d["_runs"] = runs
    return runs


def text_run_at(page, pt, tol: float = 2.0, d: dict | None = None) -> dict | None:
    """Tìm dòng chữ gốc tại điểm pt (toạ độ trang). Trả về mô tả dòng chữ để sửa."""
    pt = pymupdf.Point(pt)
    if d is None:
        d = text_dict(page)
    best = None
    for r in text_runs(d):
        if pt in r["bbox"] + (-tol, -tol, tol, tol):
            if best is None or r["bbox"].get_area() < best["bbox"].get_area():
                best = r
    return best


def text_runs_in(page, rect, d: dict | None = None) -> list[dict]:
    """Các dòng chữ nằm trong khung rect (phần lớn dòng nằm trong khung), xếp trên xuống rồi trái sang phải."""
    rect = pymupdf.Rect(rect)
    if d is None:
        d = text_dict(page)
    out = []
    for r in text_runs(d):
        b = r["bbox"]
        if b.get_area() > 0 and (b & rect).get_area() >= 0.5 * b.get_area():
            out.append(r)
    out.sort(key=lambda r: (round(r["origin"].y / max(1.0, 0.5 * r["size"])), r["bbox"].x0))
    return out


def replace_text_runs(page, runs: list[dict], lines: list, size: float | None = None) -> str:
    """Sửa cả vùng: dòng thứ k trong `lines` thay cho dòng chữ thứ k trong `runs` (giữ đúng vị trí, cỡ của dòng đó).
    Mỗi dòng là chuỗi thường hoặc danh sách mẩu nhiều kiểu. Ít dòng hơn thì các dòng thừa bị xoá; nhiều dòng hơn thì viết
    tiếp xuống dưới dòng cuối. Dòng không đổi gì thì để nguyên chữ gốc. Dòng dài quá thì tự xuống dòng."""
    used = ""
    as_segs = lambda run, x: plain_segs(run, x) if isinstance(x, str) else [tuple(v) for v in x]
    if runs and len(lines) > len(runs):
        last = runs[-1]
        ys = sorted({round(r["origin"].y, 1) for r in runs})
        pitch = (ys[-1] - ys[0]) / (len(ys) - 1) if len(ys) > 1 else 1.25 * last["size"]
        pitch = max(pitch, 1.05 * (size or last["size"]))
        for n, line in enumerate(lines[len(runs):], 1):
            segs = as_segs(last, line)
            if not segs_text(segs).strip():
                continue
            pieces = rich_lines(last, segs, size or last["size"], None)[0]
            used = write_pieces(page, pymupdf.Point(last["origin"].x, last["origin"].y + n * pitch), pieces,
                                size or last["size"]) or used
    for k in range(len(runs) - 1, -1, -1):       # từ dưới lên: đẩy nội dung bên dưới không làm lệch các dòng phía trên
        run = runs[k]
        segs = as_segs(run, lines[k]) if k < len(lines) else []
        if size is None and _norm_segs(run.get("segs", [])) == _norm_segs(segs):
            continue                             # dòng không đổi gì: để nguyên chữ gốc
        name, _n = replace_text_run_wrapped(page, run, segs, size)
        used = used or name
    return used


def write_display_text(page, disp_pt, text: str, font, size: float, color) -> None:
    """Viết một dòng chữ tại điểm (toạ độ đang hiển thị), chữ luôn nằm ngang trên màn hình dù trang bị xoay."""
    pu = pymupdf.Point(disp_pt) * page.derotation_matrix
    tw = pymupdf.TextWriter(page.rect, color=color)
    tw.append(pu, text, font=font, fontsize=size)
    if page.rotation:
        tw.write_text(page, overlay=True, morph=(pu, pymupdf.Matrix(page.rotation)))
    else:
        tw.write_text(page, overlay=True)


def insert_text_lines(page, disp_pt, text: str, size=12, color=(0, 0, 0), family="arial", bold=False,
                      italic=False) -> str:
    """Thêm chữ mới vào trang; disp_pt là góc trên trái (toạ độ hiển thị). Mỗi dòng sửa lại được về sau."""
    font = pick_font(family, 0, text, bold=bold, italic=italic)
    p = pymupdf.Point(disp_pt)
    for k, line in enumerate(text.split("\n")):
        if line.strip():
            write_display_text(page, (p.x, p.y + size * 0.9 + k * size * 1.3), line, font, size, color)
    return font.name


def replace_text_run(page, run: dict, new_text: str, size: float | None = None, color=None,
                     bold=None, italic=None, origin=None) -> str:
    """Xoá đoạn chữ gốc và viết chữ mới tại đúng vị trí. Trả về tên phông đã dùng."""
    size = size or run["size"]
    color = color if color is not None else run["color"]
    bbox = pymupdf.Rect(run["bbox"])
    dx, dy = run.get("dir", (1, 0))
    horizontal = abs(dy) < 0.01 and dx > 0
    if horizontal and bbox.width > 2:
        base = run["origin"].y
        band = pymupdf.Rect(bbox.x0 + 0.4, base - 0.45 * run["size"], bbox.x1 - 0.4, base - 0.15 * run["size"])
    else:
        band = bbox + (0.6, 0.6, -0.6, -0.6)
    if band.is_empty:
        band = bbox
    page.add_redact_annot(band, fill=False, cross_out=False)
    page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                          text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    if not new_text.strip():
        return ""
    font = pick_font(run["font"], run["flags"], new_text, bold=bold, italic=italic)
    at = pymupdf.Point(origin) if origin is not None else run["origin"]
    tw = pymupdf.TextWriter(page.rect, color=color)
    tw.append(at, new_text, font=font, fontsize=size)
    if horizontal:
        tw.write_text(page, overlay=True)
    else:
        tw.write_text(page, overlay=True, morph=(at, pymupdf.Matrix(dx, -dy, dy, dx, 0, 0)))
    return font.name


def paragraph_runs(page, run: dict, d: dict | None = None) -> list[dict]:
    """Các dòng cùng đoạn văn với `run`: cùng khối chữ của trang, hoặc các dòng liền nhau, thẳng lề trái, cỡ chữ gần bằng."""
    if d is None:
        d = text_dict(page)
    runs = [r for r in text_runs(d) if r.get("dir", (1, 0)) == run.get("dir", (1, 0))]
    same = [r for r in runs if run.get("block") is not None and r.get("block") == run.get("block")]
    if len(same) > 1:
        return sorted(same, key=lambda r: r["origin"].y)
    # file ghi mỗi dòng thành một khối riêng: tự nhận đoạn văn theo lề trái và khoảng cách dòng
    lines = sorted(runs, key=lambda r: r["origin"].y)
    k = next((j for j, r in enumerate(lines) if r is run or r["bbox"] == run["bbox"]), None)
    if k is None:
        return [run]
    para = [lines[k]]
    for step in (-1, 1):
        j, prev = k + step, lines[k]
        while 0 <= j < len(lines):
            r = lines[j]
            gap = abs(r["origin"].y - prev["origin"].y)
            if not (0.9 <= r["size"] / max(0.1, prev["size"]) <= 1.12 and abs(r["bbox"].x0 - prev["bbox"].x0) < 1.5 * prev["size"]
                    and 0.8 * prev["size"] <= gap <= 2.0 * prev["size"]):
                break
            para.append(r)
            prev, j = r, j + step
    return sorted(para, key=lambda r: r["origin"].y)


def move_text_runs(page, runs: list[dict], dx: float, dy: float) -> str:
    """Dời các dòng chữ gốc đi (dx, dy) điểm: xoá chữ ở chỗ cũ rồi viết lại đúng kiểu chữ (đậm, nghiêng, màu) ở chỗ mới."""
    plans = []
    for r in runs:
        pieces = rich_lines(r, list(r["segs"]), r["size"], None)[0]
        plans.append((r, pieces))
    for r, _ in plans:
        replace_text_run(page, r, "")
    used = ""
    for r, pieces in plans:
        used = write_pieces(page, r["origin"] + (dx, dy), pieces, r["size"]) or used
    return used


class NoRoom(Exception):
    """Không tự xuống dòng được (hết chỗ trống cuối trang, vướng ghi chú...). Thông báo nói rõ lý do."""


def wrap_words(font, size: float, text: str, width: float) -> list[str]:
    """Ngắt chữ thành các dòng không rộng quá `width` (theo đúng bề rộng chữ của phông sẽ dùng)."""
    lines, cur = [], ""
    for word in text.split():
        trial = (cur + " " + word).strip()
        if cur and font.text_length(trial, fontsize=size) > width:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    if cur:
        lines.append(cur)
    return lines or [""]


def line_frame(page, run: dict, d: dict | None = None) -> tuple[float, float]:
    """(mép phải, khoảng cách dòng) của cột chữ chứa dòng `run`: lấy theo các dòng cùng đoạn / cùng cột ở ngay trên dưới,
    để chữ sửa dài ra thì xuống dòng đúng bề rộng và giãn dòng của bản gốc."""
    if d is None:
        d = text_dict(page)
    b, size, base = run["bbox"], run["size"], run["origin"].y
    mates = []
    for r in text_runs(d):
        rb = r["bbox"]
        if r is run or rb == b or not 0.85 <= r["size"] / size <= 1.18:
            continue
        ov = min(b.x1, rb.x1) - max(b.x0, rb.x0)
        if ov > 0.5 * min(b.width, rb.width) and abs(r["origin"].y - base) <= 14 * size:
            mates.append(r)
    near = sorted(abs(r["origin"].y - base) for r in mates if 0.9 * size <= abs(r["origin"].y - base) <= 1.9 * size)
    if near:
        pitch = near[0]
    else:        # dòng đứng một mình: lấy khoảng cách dòng phổ biến của chữ cùng cỡ trên trang
        same = sorted((r for r in text_runs(d) if 0.95 <= r["size"] / size <= 1.05), key=lambda r: (round(r["bbox"].x0), r["origin"].y))
        gaps = sorted(round(q["origin"].y - p["origin"].y, 1) for p, q in zip(same, same[1:])
                      if abs(p["bbox"].x0 - q["bbox"].x0) < 1 and 1.0 * size <= q["origin"].y - p["origin"].y <= 1.9 * size)
        pitch = gaps[len(gaps) // 2] if gaps else 1.3 * size
    # chỉ tính các dòng liền mạch với dòng này (cách nhau đều theo khoảng cách dòng), không lấy nhầm đoạn văn khác
    chain = {round(base, 1)}
    for sign in (1, -1):
        y = base
        for r in sorted((r for r in mates if (r["origin"].y - base) * sign > 0), key=lambda r: (r["origin"].y - base) * sign):
            if abs(abs(r["origin"].y - y) - pitch) < 0.3 * pitch:
                y = r["origin"].y
                chain.add(round(y, 1))
            elif abs(r["origin"].y - y) > 1.3 * pitch:
                break
    rights = [r["bbox"].x1 for r in mates if round(r["origin"].y, 1) in chain]
    if rights:
        return max(rights + [b.x1]), pitch
    # dòng đứng một mình (tiêu đề, một ô của bảng): được dài tới vật cản gần nhất bên phải, hoặc tới lề phải của trang
    left = min((r["bbox"].x0 for r in text_runs(d)), default=b.x0)
    limit = page.rect.width - max(18.0, left)
    for r in text_runs(d):
        rb = r["bbox"]
        if rb.x0 > b.x1 and min(b.y1, rb.y1) - max(b.y0, rb.y0) > 0:
            limit = min(limit, rb.x0 - 0.8 * size)
    try:
        for g in page.get_drawings():
            gr = pymupdf.Rect(g["rect"])
            if gr.x0 > b.x1 - 1 and gr.width < 3 and gr.y0 < b.y1 and gr.y1 > b.y0:     # nét kẻ dọc (viền ô bảng)
                limit = min(limit, gr.x0 - 3)
    except Exception:
        pass
    return max(limit, b.x1), pitch


def shift_below(page, prev_bottom: float, delta: float, d: dict | None = None, cut: float | None = None) -> None:
    """Đẩy mọi thứ nằm dưới mức `prev_bottom` của trang xuống `delta` điểm để lấy chỗ cho dòng chữ mới
    (chữ, hình, bảng, nét kẻ bên dưới giữ nguyên dạng; nét kẻ dọc và nền ô bảng được kéo dài qua khoảng trống mới)."""
    doc = page.parent
    if page.rotation:
        raise NoRoom("Trang đang bị xoay nên app không tự đẩy nội dung xuống được.")
    W, H = page.rect.width, page.rect.height
    if d is None:
        d = text_dict(page)
    runs = text_runs(d)
    below = [r for r in runs if r["bbox"].y0 >= prev_bottom - 0.35 * r["size"] and r["origin"].y > prev_bottom]
    next_top = min((r["bbox"].y0 for r in below), default=H)
    if cut is None:
        cut = prev_bottom + 0.6 if next_top - prev_bottom >= 1.2 else (prev_bottom + max(prev_bottom, next_top)) / 2
    else:                                        # cắt tại mức cho sẵn (lề dưới của một hàng trong bảng: không có chữ ở đó)
        next_top = cut + 0.4
        prev_bottom = cut - 0.4
    # Phần chân trang (số trang, dòng chân trang) đứng yên; chỉ phần thân bên dưới dòng đang sửa mới bị đẩy xuống.
    items = []
    for kind, bb in page.get_bboxlog():
        r = pymupdf.Rect(bb)
        if r.y1 > cut + 0.5 and r.get_area() < 0.7 * W * H and r.y1 <= H + 1:
            items.append(r)
    foot = [r for r in items if r.y0 > 0.88 * H]
    body = [r for r in items if r.y0 <= 0.88 * H]
    lowest = max((r.y1 for r in body), default=cut)
    floor = H - 2.0                              # mức thấp nhất mà phần thân được phép chạm tới
    if foot and min(r.y0 for r in foot) - lowest >= 10:
        floor = min(r.y0 for r in foot) - 3.0
    else:
        lowest = max([lowest] + [r.y1 for r in foot])
    if lowest + delta > floor - 1:
        raise NoRoom(f"Cuối trang không còn đủ chỗ trống để đẩy phần bên dưới xuống (cần thêm khoảng {delta:.0f} pt, "
                     f"chỉ còn {max(0.0, floor - 1 - lowest):.0f} pt). Hãy rút ngắn câu hoặc giảm cỡ chữ.")
    if lowest <= cut + 0.5:
        return                                   # bên dưới không có gì: không cần đẩy
    has_foot = floor < H - 2.5
    fixed = ("Highlight", "Underline", "StrikeOut", "Squiggly", "Ink", "Line", "Polygon", "PolyLine", "Redact")
    move = []
    for an in page.annots() or []:
        if cut < an.rect.y1 and an.rect.y0 < floor:
            if an.type[1] in fixed or an.rect.y0 < cut:
                raise NoRoom("Bên dưới dòng này có ghi chú / tô sáng nên app không tự đẩy nội dung xuống được. "
                             "Hãy làm phẳng ghi chú (Sửa → Làm phẳng ghi chú) hoặc xoá chúng trước.")
            move.append(an.xref)
    if any(cut < w.rect.y1 and w.rect.y0 < floor for w in (page.widgets() or [])):
        raise NoRoom("Bên dưới dòng này có ô biểu mẫu / khung chữ ký nên app không tự đẩy nội dung xuống được.")
    crossing = any(pymupdf.Rect(im["bbox"]).y0 < y < pymupdf.Rect(im["bbox"]).y1
                   for im in page.get_image_info() for y in (cut, floor))
    img_mode = pymupdf.PDF_REDACT_IMAGE_NONE if crossing else pymupdf.PDF_REDACT_IMAGE_REMOVE
    z = max(1, round(delta * 0.8)) / delta       # tỉ lệ dựng ảnh kiểm tra: chọn sao cho delta đúng bằng số nguyên điểm ảnh
    before = page.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False, annots=False)

    # Bốn bản sao của trang trong một tài liệu tạm, mỗi bản chỉ giữ chữ của một phần: (0) phần trên, (1) phần thân bên dưới,
    # (2) không còn chữ (để kéo dài nét kẻ dọc / nền ô bảng qua khoảng trống mới mở ra), (3) phần chân trang.
    # Chữ phải được gỡ thật khỏi từng bản (không chỉ che đi), nếu không khi tìm / chọn chữ sẽ thấy mỗi dòng hai lần.
    tmp = pymupdf.open()
    for _ in range(4 if has_foot else 3):
        tmp.insert_pdf(doc, from_page=page.number, to_page=page.number, annots=False, links=False, widgets=False)
    full = pymupdf.Rect(0, 0, W, H)
    plans = [(0, [pymupdf.Rect(0, prev_bottom + 0.3, W, H)]),
             (1, [pymupdf.Rect(0, 0, W, next_top - 0.3)] + ([pymupdf.Rect(0, floor, W, H)] if has_foot else [])),
             (2, [full])] + ([(3, [pymupdf.Rect(0, 0, W, floor)])] if has_foot else [])
    # nét kẻ / hình chỉ được gỡ khi nằm hẳn về phía bên kia của mức cắt (khung gỡ chữ thì rộng hơn một chút để bắt trọn dòng)
    gfx = {0: [pymupdf.Rect(0, cut + 0.2, W, H)],
           1: [pymupdf.Rect(0, 0, W, cut - 0.2)] + ([pymupdf.Rect(0, floor, W, H)] if has_foot else []),
           3: [pymupdf.Rect(0, 0, W, floor)]}
    for k, rects in plans:
        tp = tmp[k]
        for rect in rects:
            if not rect.is_empty:
                tp.add_redact_annot(rect, fill=False, cross_out=False)
        tp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                            text=pymupdf.PDF_REDACT_TEXT_REMOVE)
        for rect in gfx.get(k, []):
            if not rect.is_empty:
                tp.add_redact_annot(rect, fill=False, cross_out=False)
        if k in gfx:
            tp.apply_redactions(images=img_mode, graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                                text=pymupdf.PDF_REDACT_TEXT_NONE)
    doc.xref_set_key(page.xref, "Contents", "null")
    doc.xref_set_key(page.xref, "Resources", "<<>>")
    top, mid = pymupdf.Rect(0, 0, W, cut), pymupdf.Rect(0, cut, W, floor - delta)
    page.show_pdf_page(top, tmp, 0, clip=top)
    strip = pymupdf.Rect(0, cut - 0.25, W, cut + 0.25)
    page.show_pdf_page(pymupdf.Rect(0, cut, W, cut + delta), tmp, 2, clip=strip, keep_proportion=False)
    page.show_pdf_page(mid + (0, delta, 0, delta), tmp, 1, clip=mid)
    if has_foot:
        tail = pymupdf.Rect(0, floor, W, H)
        page.show_pdf_page(tail, tmp, 3, clip=tail)
    tmp.close()
    for xref in move:
        an = page.load_annot(xref)
        an.set_rect(an.rect + (0, delta, 0, delta))
        an.update()
    for ln in page.get_links():
        if cut - 1 <= ln["from"].y0 < floor:
            ln["from"] = ln["from"] + (0, delta, 0, delta)
            page.update_link(ln)
    # kiểm tra bằng mắt máy: phần trên và chân trang phải y nguyên, phần thân phải y nguyên nhưng thấp hơn đúng delta
    after = page.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False, annots=False)
    n, stride = before.n, before.stride
    a, b2 = before.samples, after.samples

    def differ(y_a0: int, y_b0: int, rows: int) -> float:
        bad = total = 0
        for k in range(0, max(0, rows), 2):
            ra = a[(y_a0 + k) * stride:(y_a0 + k) * stride + before.width * n]
            rb = b2[(y_b0 + k) * stride:(y_b0 + k) * stride + after.width * n]
            total += 1
            if ra != rb and sum(abs(x - y) for x, y in zip(ra[::7], rb[::7])) > 0.02 * 255 * len(ra[::7]):
                bad += 1
        return bad / max(1, total)
    yc, yf = int(cut * z), int(floor * z)
    sh = int(round(delta * z))                       # dịch theo điểm ảnh (z đã chọn để số này là số nguyên)
    if differ(0, 0, yc - 2) > 0.02 or differ(yc + 3, yc + 3 + sh, yf - sh - yc - 8) > 0.05 \
            or differ(yf + 3, yf + 3, before.height - yf - 5) > 0.02:
        raise NoRoom("Không tự đẩy được phần bên dưới ở vị trí này (có chữ nằm ngang hàng ở cột bên cạnh, hoặc bố cục "
                     "quá phức tạp). Hãy rút ngắn câu, hoặc dùng Xoá trắng một vùng + Thêm chữ.")


def plain_segs(run: dict, text: str, color=None, bold=None, italic=None) -> list[tuple]:
    """Một dòng chữ thường (một kiểu) dưới dạng danh sách mẩu: [(chữ, đậm, nghiêng, màu)]."""
    return [(text, bool(run["flags"] & 16) if bold is None else bool(bold),
             bool(run["flags"] & 2) if italic is None else bool(italic), color if color is not None else run["color"])]


def _norm_segs(segs) -> list[tuple]:
    """Dạng chuẩn để so hai dòng nhiều kiểu chữ: gộp mẩu liền nhau cùng kiểu, bỏ khoảng trắng đầu/cuối, làm tròn màu."""
    out: list[list] = []
    for text, b, i, col in segs:
        key = [bool(b), bool(i), tuple(round(float(v) * 255) for v in col)]
        if out and out[-1][1:] == key:
            out[-1][0] += text
        elif text:
            out.append([text] + key)
    if out:
        out[0][0] = out[0][0].lstrip()
        out[-1][0] = out[-1][0].rstrip()
    return [tuple(x) for x in out if x[0]]


def segs_text(segs) -> str:
    return "".join(x[0] for x in segs)


def rich_lines(run: dict, segs: list[tuple], size: float, width: float | None) -> list[list[tuple]]:
    """Dàn các mẩu chữ nhiều kiểu thành các dòng không rộng quá `width` (None = một dòng).
    Mỗi dòng: [(chữ, phông, màu), ...] viết nối tiếp nhau từ trái sang."""
    fonts: dict = {}

    def font_of(b, i, text):
        key = (b, i)
        if key not in fonts or not all(fonts[key].has_glyph(ord(c)) for c in text if not c.isspace()):
            fonts[key] = pick_font(run["font"], run["flags"], text, bold=b, italic=i)
        return fonts[key]
    words: list[list[tuple]] = []        # mỗi từ = các mẩu dính liền nhau (một từ có thể nửa đậm nửa thường)
    cur: list[tuple] = []
    for text, b, i, col in segs:
        for part in re.split(r"(\s+)", text):
            if not part:
                continue
            if part.isspace():
                if cur:
                    words.append(cur)
                    cur = []
            else:
                cur.append((part, font_of(b, i, part), col, b, i))
    if cur:
        words.append(cur)
    lines: list[list[tuple]] = [[]]
    x = 0.0
    space = font_of(False, False, " ").text_length(" ", fontsize=size)
    for w in words:
        wlen = sum(f.text_length(t, fontsize=size) for t, f, _c, _b, _i in w)
        if lines[-1] and width is not None and x + space + wlen > width + 0.6:
            lines.append([])
            x = 0.0
        if lines[-1]:
            last = lines[-1][-1]
            # khoảng trắng mang kiểu của mẩu đứng trước nó
            lines[-1][-1] = (last[0] + " ", last[1], last[2])
            x += last[1].text_length(" ", fontsize=size)
        for t, f, c, _b, _i in w:
            if lines[-1] and lines[-1][-1][1] is f and lines[-1][-1][2] == c:
                lines[-1][-1] = (lines[-1][-1][0] + t, f, c)
            else:
                lines[-1].append((t, f, c))
        x += wlen
    return [ln for ln in lines if ln] or [[]]


def write_pieces(page, origin, pieces: list[tuple], size: float) -> str:
    """Viết một dòng gồm nhiều mẩu (chữ, phông, màu) nối tiếp nhau. Trả về tên phông đầu tiên đã dùng."""
    x, y = origin.x, origin.y
    by_color: dict = {}
    name = ""
    for text, font, col in pieces:
        tw = by_color.setdefault(tuple(col), pymupdf.TextWriter(page.rect, color=col))
        tw.append(pymupdf.Point(x, y), text, font=font, fontsize=size)
        x += font.text_length(text, fontsize=size)
        name = name or font.name
    for tw in by_color.values():
        tw.write_text(page, overlay=True)
    return name


def cell_of(page, bbox) -> dict | None:
    """Ô bảng chứa dòng chữ (dựa vào các nét kẻ quanh nó): {"x0","x1","top","bottom"} hoặc None nếu không nằm trong bảng."""
    b = pymupdf.Rect(bbox)
    W = page.rect.width
    vs, hs = [], []
    try:
        for g in page.get_drawings():
            r = pymupdf.Rect(g["rect"])
            if r.width <= 2.5 and r.height >= 6:
                vs.append(r)
            elif r.height <= 2.5 and r.width >= 6:
                hs.append(r)
            elif r.width > 6 and r.height > 6 and g.get("type") in ("s", "fs"):      # khung chữ nhật có viền: bốn cạnh là nét kẻ
                vs += [pymupdf.Rect(r.x0, r.y0, r.x0, r.y1), pymupdf.Rect(r.x1, r.y0, r.x1, r.y1)]
                hs += [pymupdf.Rect(r.x0, r.y0, r.x1, r.y0), pymupdf.Rect(r.x0, r.y1, r.x1, r.y1)]
    except Exception:
        return None
    left = [r.x1 for r in vs if r.x1 <= b.x0 + 1 and r.y0 <= b.y0 + 2 and r.y1 >= b.y1 - 2]
    right = [r.x0 for r in vs if r.x0 >= b.x1 - 1 and r.y0 <= b.y0 + 2 and r.y1 >= b.y1 - 2]
    if not left and not right:
        return None
    x0, x1 = (max(left) if left else 0.0), (min(right) if right else W)
    span = lambda r: min(r.x1, x1) - max(r.x0, x0) >= 0.8 * (x1 - x0)
    below = [r.y0 for r in hs if r.y0 >= b.y1 - 1 and span(r)]
    above = [r.y1 for r in hs if r.y1 <= b.y0 + 1 and span(r)]
    if not below:
        return None
    return {"x0": x0, "x1": x1, "bottom": min(below), "top": max(above) if above else None}


def _same_rows(a, b, x0: int, x1: int, ya: int, yb: int, rows: int) -> float:
    """Tỉ lệ số hàng điểm ảnh khác nhau giữa hai ảnh (cột x0..x1; hàng ya.. của a so với yb.. của b)."""
    n = a.n
    sa, sb = a.samples, b.samples
    bad = total = 0
    for k in range(0, max(0, rows), 2):
        ra = sa[(ya + k) * a.stride + x0 * n:(ya + k) * a.stride + x1 * n]
        rb = sb[(yb + k) * b.stride + x0 * n:(yb + k) * b.stride + x1 * n]
        total += 1
        if ra != rb and sum(abs(p - q) for p, q in zip(ra[::7], rb[::7])) > 0.02 * 255 * max(1, len(ra[::7])):
            bad += 1
    return bad / max(1, total)


def shift_in_box(page, x0: float, x1: float, cut: float, floor: float, delta: float) -> None:
    """Đẩy riêng phần nằm trong khung x0..x1, từ mức `cut` tới `floor`, xuống `delta` điểm; mọi thứ ngoài khung đứng yên.
    Dùng cho ô bảng: chỉ các dòng còn lại của chính ô đó bị đẩy, ô bên cạnh không bị đụng tới."""
    doc = page.parent
    W, H = page.rect.width, page.rect.height
    box = pymupdf.Rect(x0, cut, x1, floor)
    runs = text_runs(text_dict(page))
    inside = [r for r in runs if r["bbox"].intersects(box) and r["origin"].y > cut]
    if not inside:
        return
    if any(r["bbox"].x0 < x0 - 0.5 or r["bbox"].x1 > x1 + 0.5 for r in inside):
        raise NoRoom("Chữ ở ô này tràn sang ô bên cạnh nên app không đẩy riêng trong ô được.")
    if max(r["bbox"].y1 for r in inside) + delta > floor + 0.5:
        raise NoRoom("Ô này không còn đủ chỗ để đẩy các dòng bên dưới xuống.")
    if any(an.rect.intersects(box) for an in (page.annots() or [])) or any(w.rect.intersects(box) for w in (page.widgets() or [])):
        raise NoRoom("Trong ô này có ghi chú / ô biểu mẫu nên app không tự đẩy nội dung xuống được.")
    next_top = min(r["bbox"].y0 for r in inside)
    z = max(1, round(delta * 0.8)) / delta
    before = page.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False, annots=False)
    tmp = pymupdf.open()
    for _ in range(2):
        tmp.insert_pdf(doc, from_page=page.number, to_page=page.number, annots=False, links=False, widgets=False)
    # bản 0 (đứng yên): gỡ chữ của phần sẽ bị đẩy; bản 1 (bị đẩy): chỉ giữ chữ của phần đó
    stay, move = tmp[0], tmp[1]
    stay.add_redact_annot(pymupdf.Rect(x0, next_top + 0.3, x1, floor), fill=False, cross_out=False)
    stay.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_COVERED,
                          text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    for rect in (pymupdf.Rect(0, 0, W, next_top - 0.3), pymupdf.Rect(0, floor, W, H),
                 pymupdf.Rect(0, 0, x0, H), pymupdf.Rect(x1, 0, W, H)):
        if not rect.is_empty:
            move.add_redact_annot(rect, fill=False, cross_out=False)
    move.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                          text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    doc.xref_set_key(page.xref, "Contents", "null")
    doc.xref_set_key(page.xref, "Resources", "<<>>")
    page.show_pdf_page(page.rect, tmp, 0)
    src = pymupdf.Rect(x0, cut, x1, floor - delta)
    page.show_pdf_page(src + (0, delta, 0, delta), tmp, 1, clip=src)
    tmp.close()
    after = page.get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False, annots=False)
    xa, xb, yc, yf, sh = int(x0 * z) + 2, int(x1 * z) - 2, int(cut * z), int(floor * z), int(round(delta * z))
    wpx = before.width
    if _same_rows(before, after, 0, wpx, 0, 0, yc - 2) > 0.02 or _same_rows(before, after, 0, wpx, yf + 3, yf + 3, before.height - yf - 5) > 0.02 \
            or _same_rows(before, after, 0, max(0, xa - 4), yc, yc, yf - yc) > 0.02 \
            or _same_rows(before, after, min(wpx, xb + 4), wpx, yc, yc, yf - yc) > 0.02 \
            or _same_rows(before, after, xa, xb, yc + 3, yc + 3 + sh, yf - sh - yc - 8) > 0.06:
        raise NoRoom("Không đẩy riêng được các dòng trong ô này (bố cục quá phức tạp).")


def replace_text_run_wrapped(page, run: dict, new, size: float | None = None, color=None,
                             bold=None, italic=None) -> tuple[str, int]:
    """Thay một dòng chữ gốc bằng chữ mới. `new` là chuỗi thường, hoặc danh sách mẩu nhiều kiểu [(chữ, đậm, nghiêng, màu)].
    Chữ mới dài quá bề rộng của cột / ô thì tự xuống dòng (đúng khoảng cách dòng của đoạn gốc):
      - dòng nằm trong ô bảng: chỉ các dòng còn lại của chính ô đó bị đẩy xuống; ô hết chỗ thì cả hàng của bảng cao thêm
        (phần bên dưới bảng lùi xuống), chữ ở các ô khác cùng hàng đứng yên;
      - dòng thường: phần bên dưới của trang được đẩy xuống.
    Trả về (tên phông, số dòng đã viết). Không đẩy được thì báo NoRoom, trang giữ nguyên (do model.edit hoàn lại)."""
    segs = plain_segs(run, new, color, bold, italic) if isinstance(new, str) else [tuple(x) for x in new]
    text = segs_text(segs)
    dx, dy = run.get("dir", (1, 0))
    if not text.strip() or not (abs(dy) < 0.01 and dx > 0):
        s0 = segs[0] if segs else (text, None, None, None)
        return replace_text_run(page, run, text, size, s0[3], bold=s0[1], italic=s0[2]), 1
    fsize = size or run["size"]
    d = text_dict(page)
    me = next((r for r in text_runs(d) if r["bbox"] == run["bbox"]), run)
    limit, pitch = line_frame(page, me, d)
    pitch = max(pitch, 1.05 * fsize) if size else pitch
    cell = cell_of(page, run["bbox"])
    if cell:                                     # trong ô bảng: được dài tới sát viền phải của ô (trừ lề trong của ô)
        pad = min(12.0, max(2.0, run["bbox"].x0 - cell["x0"]))
        limit = max(limit if limit <= cell["x1"] - 1 else 0, cell["x1"] - pad)
    room = limit - run["origin"].x
    lines = rich_lines(run, segs, fsize, room if room >= 4 * fsize else None)
    replace_text_run(page, run, "")              # gỡ dòng chữ cũ
    extra = (len(lines) - 1) * pitch
    if extra > 0:
        bottom = pymupdf.Rect(run["bbox"]).y1
        if cell:
            runs = text_runs(text_dict(page))
            mine = [r for r in runs if r["bbox"].x0 >= cell["x0"] - 0.5 and r["bbox"].x1 <= cell["x1"] + 0.5
                    and (cell["top"] or 0) - 1 <= r["bbox"].y0 and r["bbox"].y1 <= cell["bottom"] + 1]
            lowest = max([r["bbox"].y1 for r in mine] + [bottom])
            pad_top = min(12.0, max(2.0, (min(r["bbox"].y0 for r in mine) if mine else run["bbox"].y0) - cell["top"])) \
                if cell["top"] is not None else 3.0
            grow = lowest + extra + pad_top - cell["bottom"]
            if grow > 0.5:                       # ô hết chỗ: cả hàng cao thêm, phần bên dưới bảng lùi xuống
                shift_below(page, cell["bottom"] - 1.0, grow, cut=cell["bottom"] - 1.0)
            if any(r["origin"].y > bottom for r in mine):       # các dòng còn lại của chính ô này: đẩy riêng trong ô
                shift_in_box(page, cell["x0"] + 0.8, cell["x1"] - 0.8, bottom + 0.4,
                             cell["bottom"] + max(0.0, grow) - 1.2, extra)
        else:
            shift_below(page, bottom, extra)
    name = ""
    for k, pieces in enumerate(lines):
        name = write_pieces(page, pymupdf.Point(run["origin"].x, run["origin"].y + k * pitch), pieces, fsize) or name
    return name, len(lines)


def finalize_fonts(doc) -> None:
    """Thu gọn phông đã nhúng (chỉ giữ ký tự đang dùng) để file không bị phình."""
    try:
        doc.subset_fonts()
    except Exception:
        pass


# ------------------------------------------------------------ thông tin tài liệu
def pdf_date(v: str | None) -> str:
    """'D:20261010113000+07'00'' -> '10/10/2026 11:30:00 (+07:00)'."""
    if not v:
        return ""
    m = re.match(r"(?:D:)?(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?\s*([+\-Zz])?(\d{2})?'?(\d{2})?", v)
    if not m:
        return v
    y, mo, d, h, mi, sec, sign, th, tm = m.groups()
    out = f"{d or '01'}/{mo or '01'}/{y}"
    if h:
        out += f" {h}:{mi or '00'}:{sec or '00'}"
    if sign in ("+", "-") and th:
        out += f" ({sign}{th}:{tm or '00'})"
    elif sign in ("Z", "z"):
        out += " (UTC)"
    return out


_PAPER = [("A5", 148, 210), ("A4", 210, 297), ("A3", 297, 420), ("A2", 420, 594), ("B5", 176, 250),
          ("Letter", 216, 279), ("Legal", 216, 356)]


def page_size_text(w_pt: float, h_pt: float) -> str:
    w, h = w_pt * 25.4 / 72, h_pt * 25.4 / 72
    name = ""
    for n, a, b in _PAPER:
        if abs(min(w, h) - a) < 2.5 and abs(max(w, h) - b) < 2.5:
            name = f"{n} {'ngang' if w > h else 'dọc'}, "
            break
    return f"{name}{w:.0f} × {h:.0f} mm ({w_pt:.0f} × {h_pt:.0f} pt)"


def security_info(data: bytes, password: str | None) -> dict:
    """Kiểu mã hoá và các quyền ghi trong file gốc (đọc trước khi giải mã trong bộ nhớ)."""
    try:
        d = pymupdf.open(stream=data, filetype="pdf")
        if d.needs_pass:
            d.authenticate(password or "")
        perm = d.permissions
        out = {"method": (d.metadata or {}).get("encryption") or "Có",
               "print": bool(perm & pymupdf.PDF_PERM_PRINT), "copy": bool(perm & pymupdf.PDF_PERM_COPY),
               "modify": bool(perm & pymupdf.PDF_PERM_MODIFY), "annotate": bool(perm & pymupdf.PDF_PERM_ANNOTATE),
               "form": bool(perm & pymupdf.PDF_PERM_FORM)}
        d.close()
        return out
    except Exception:
        return {}


def document_report(doc, data: bytes | None, current: int = 0, max_pages: int = 300) -> dict:
    """Gom thông tin chi tiết của tài liệu để hiển thị."""
    n = doc.page_count
    head = (data or b"")[:1024]
    m = re.search(rb"%PDF-(\d\.\d)", head)
    version = m.group(1).decode() if m else ""
    try:
        cat = doc.pdf_catalog()
        v = doc.xref_get_key(cat, "Version")
        if v[0] == "name" and v[1].lstrip("/") > version:
            version = v[1].lstrip("/")
        tagged = doc.xref_get_key(cat, "MarkInfo")[0] != "null"
    except Exception:
        tagged = False
    sizes: dict[tuple, int] = {}
    rotated = 0
    for i in range(n):
        try:
            r = doc.page_cropbox(i) if hasattr(doc, "page_cropbox") else doc[i].rect
            rot = doc[i].rotation if i < max_pages else 0
        except Exception:
            continue
        w, h = (r.height, r.width) if rot in (90, 270) else (r.width, r.height)
        key = (round(w), round(h))
        sizes[key] = sizes.get(key, 0) + 1
        rotated += 1 if rot else 0
    top = sorted(sizes.items(), key=lambda kv: -kv[1])
    if len(top) == 1:
        size_text = page_size_text(*top[0][0])
    else:
        size_text = "; ".join(f"{page_size_text(*k)}: {c} trang" for k, c in top[:4])
        if len(top) > 4:
            size_text += f"; và {len(top) - 4} khổ khác"
    cur = doc[current]
    cur_text = page_size_text(cur.rect.width, cur.rect.height) + (f", xoay {cur.rotation}°" if cur.rotation else "")
    scan = min(n, max_pages)
    fonts: dict[str, str] = {}
    images, annots, links, widgets, with_text, sigs = set(), 0, 0, 0, 0, 0
    for i in range(scan):
        page = doc[i]
        try:
            for f in doc.get_page_fonts(i):
                name = f[3].split("+")[-1]
                emb = "nhúng" if f[1] not in ("n/a", "") else "không nhúng"
                fonts.setdefault(name, f"{f[2]}, {emb}")
            for im in doc.get_page_images(i):
                images.add(im[0])
            if page.get_text("text", clip=None).strip():
                with_text += 1
            annots += sum(1 for _ in page.annots())
            links += len(page.get_links())
            for w in page.widgets():
                widgets += 1
                if w.field_type == pymupdf.PDF_WIDGET_TYPE_SIGNATURE and w.is_signed:
                    sigs += 1
        except Exception:
            continue
    if with_text == 0:
        text = "Không có (bản scan/ảnh, cần OCR để tìm kiếm và chọn chữ)"
    elif with_text == scan:
        text = "Có ở mọi trang"
    else:
        text = f"Có ở {with_text}/{scan} trang (các trang còn lại là ảnh)"
    font_html = "<br>".join(f"{k} <span style='color:#80868b'>({v})</span>" for k, v in sorted(fonts.items())[:40])
    if len(fonts) > 40:
        font_html += f"<br>… và {len(fonts) - 40} phông khác"
    if not sigs and data and b"/ByteRange" in data:
        sigs = data.count(b"/ByteRange")
    try:
        toc = len(doc.get_toc(simple=True))
    except Exception:
        toc = 0
    try:
        emb = doc.embfile_count()
    except Exception:
        emb = 0
    return {"version": version, "linearized": b"/Linearized" in head, "tagged": tagged, "sizes": size_text,
            "current": cur_text, "rotated": f"{rotated} trang" if rotated else "Không",
            "toc": toc, "text": text, "images": len(images), "annots": annots, "links": links,
            "widgets": widgets, "embedded": emb, "fonts": font_html or "Không có",
            "signatures": sigs, "scope": f" ({scan} trang đầu)" if scan < n else ""}


_REF = r"(?:[A-Za-z]{0,2}\d{1,4}(?:\s?[-–.]\s?\d{1,4})?|[ivxlcdm]{1,7})"
_LEADER = re.compile(r"^(?P<title>.+?)\s*(?:[.…·•_]\s?){4,}\s*(?P<ref>" + _REF + r")\s*$", re.I)
_REF_ONLY = re.compile(r"^" + _REF + r"$", re.I)


def _norm_text(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip().lower()


def _norm_ref(r: str) -> str:
    return re.sub(r"\s+", "", r).replace("–", "-").lower()


def _page_rows(page) -> list[dict]:
    """Gom các đoạn chữ cùng một dòng ngang (kể cả khi nằm ở các khối khác nhau) thành từng hàng."""
    spans = []
    for b in text_dict(page).get("blocks", []):
        for ln in b.get("lines", []):
            dx, dy = ln.get("dir", (1, 0))
            if abs(dy) > 0.01 or dx <= 0:
                continue
            for sp in ln["spans"]:
                if sp["text"].strip():
                    spans.append(sp)
    spans.sort(key=lambda sp: (sp["origin"][1], sp["bbox"][0]))
    rows: list[list] = []
    for sp in spans:
        if rows and abs(rows[-1][-1]["origin"][1] - sp["origin"][1]) < 2.5:
            rows[-1].append(sp)
        else:
            rows.append([sp])
    out = []
    for r in rows:
        r.sort(key=lambda sp: sp["bbox"][0])
        text = ""
        for k, sp in enumerate(r):
            if k and sp["bbox"][0] - r[k - 1]["bbox"][2] > 1.0 and not text.endswith(" "):
                text += " "
            text += sp["text"]
        main = max(r, key=lambda sp: len(sp["text"].strip(". ")))
        out.append({"text": " ".join(text.split()), "spans": r, "x0": r[0]["bbox"][0], "x1": r[-1]["bbox"][2],
                    "y0": min(sp["bbox"][1] for sp in r), "y": r[0]["origin"][1], "size": round(main["size"] * 2) / 2,
                    "bold": bool(main["flags"] & 16) or "bold" in main["font"].lower()})
    return out


def _toc_entries_on_page(page) -> list[dict]:
    """Các dòng dạng mục lục in trên trang: 'Tiêu đề ........ 12' hoặc tiêu đề + số trang canh lề phải."""
    ents = []
    for row in _page_rows(page):
        t = row["text"]
        m = _LEADER.match(t)
        title = ref = None
        if m:
            title, ref = m.group("title"), m.group("ref")
        elif len(row["spans"]) >= 2:
            last = row["spans"][-1]
            gap = last["bbox"][0] - row["spans"][-2]["bbox"][2]
            if _REF_ONLY.match(last["text"].strip()) and gap >= 30:
                ref = last["text"].strip()
                title = " ".join(" ".join(sp["text"] for sp in row["spans"][:-1]).split())
        if not title or not ref:
            continue
        title = title.strip(" .…·•_\t")
        if len(title) < 2 or not any(c.isalnum() for c in title):
            continue
        ents.append({"title": title[:150], "ref": _norm_ref(ref), "x0": row["x0"], "x1": row["x1"], "size": row["size"],
                     "bold": row["bold"], "leader": bool(m)})
    if len(ents) < 4:
        return []
    xs = sorted(e["x1"] for e in ents)
    med = xs[len(xs) // 2]
    aligned = [e for e in ents if abs(e["x1"] - med) <= 10 or e["leader"]]   # số trang canh thẳng lề phải
    return aligned if len(aligned) >= 4 else []


class _Nester:
    """Đổi cấp thô (theo thụt lề / cỡ chữ) thành cấp lồng nhau hợp lệ: mục không có mục cha phía trước
    thì được đưa lên ngang hàng, các mục cùng cấp thô luôn ngang hàng nhau."""

    def __init__(self, max_levels: int):
        self.stack: list[float] = []
        self.max = max_levels

    def level(self, raw: float) -> int:
        while self.stack and self.stack[-1] >= raw:
            self.stack.pop()
        if len(self.stack) >= self.max:      # quá số cấp cho phép: gộp vào cấp sâu nhất
            del self.stack[self.max - 1:]
        self.stack.append(raw)
        return len(self.stack)


def _levels_from_indent(ents: list[dict], max_levels: int) -> None:
    """Gán cấp theo độ thụt lề của tiêu đề trong trang mục lục (cùng lề thì chữ to/đậm hơn là cấp cao hơn)."""
    xs = sorted({round(e["x0"]) for e in ents})
    groups: list[list[int]] = []
    for x in xs:
        if groups and x - groups[-1][-1] <= 4:
            groups[-1].append(x)
        else:
            groups.append([x])
    lvl_of = {x: k + 1 for k, g in enumerate(groups) for x in g}
    for e in ents:
        e["level"] = lvl_of[round(e["x0"])]
    if len(groups) == 1:  # không thụt lề: phân cấp theo cỡ chữ / đậm
        styles = sorted({(e["size"], e["bold"]) for e in ents}, key=lambda s: (-s[0], not s[1]))
        for e in ents:
            e["level"] = styles.index((e["size"], e["bold"])) + 1
    for e in ents:
        e["raw"] = e["level"]


def _resolve_toc_pages(doc, ents: list[dict], first_body: int) -> int:
    """Tìm trang thật của từng mục (ents[i]['page']). Trả về số mục tìm được."""
    n = doc.page_count
    texts: dict[int, str] = {}

    def text_of(i: int) -> str:
        if i not in texts:
            try:
                texts[i] = _norm_text(doc[i].get_text("text"))
            except Exception:
                texts[i] = ""
        return texts[i]

    # 1) nhãn trang của chính file PDF, 2) số trang in ở đầu/cuối mỗi trang
    label_map: dict[str, int] = {}
    nums: dict[int, set] = {}
    try:
        has_labels = bool(doc.get_page_labels())
    except Exception:
        has_labels = False
    limit = min(n, 3000)
    for i in range(limit):
        page = doc[i]
        if has_labels:
            try:
                label_map.setdefault(_norm_ref(page.get_label()), i)
            except Exception:
                pass
        h = page.rect.height
        toks = set()
        try:
            for w in page.get_text("words"):
                if w[3] < h * 0.13 or w[1] > h * 0.87:
                    tk = _norm_ref(w[4].strip("()[]|.,"))
                    if tk:
                        toks.add(tk)
        except Exception:
            pass
        nums[i] = toks
        if i >= first_body:
            for tk in toks:
                if _REF_ONLY.match(tk) and not tk.isdigit():
                    label_map.setdefault(tk, i)
    numeric = [e for e in ents if e["ref"].isdigit()]
    offset = None
    if numeric:
        best, best_d = 0, None
        for d in range(-12, 120):
            sc = sum(1 for e in numeric if 0 <= int(e["ref"]) + d < limit and e["ref"] in nums.get(int(e["ref"]) + d, ()))
            if sc > best:
                best, best_d = sc, d
        if best >= max(2, len(numeric) * 0.3):
            offset = best_d
    found, prev = 0, first_body
    for e in ents:
        key = _norm_text(e["title"])[:60]
        cand = None
        if e["ref"].isdigit():
            if offset is not None and 0 <= int(e["ref"]) + offset < n:
                cand = int(e["ref"]) + offset
            elif e["ref"] in label_map and has_labels:
                cand = label_map[e["ref"]]
        else:
            cand = label_map.get(e["ref"])
        page = None
        if cand is not None:
            for c in (cand, cand + 1, cand - 1):
                if 0 <= c < n and key and key in text_of(c):
                    page = c
                    break
        if page is None and key:   # dò tiêu đề trong nội dung, tính từ mục trước trở đi
            for c in range(max(prev, first_body), min(n, max(prev, first_body) + 400)):
                if key in text_of(c):
                    page = c
                    break
        if page is None and cand is not None:
            page = cand
        if page is None:
            e["page"] = None
            continue
        e["page"] = page
        prev = max(prev, page) if page >= prev else prev
        found += 1
    return found


def headings_from_printed_toc(doc, max_levels: int = 4, scan_pages: int = 40) -> tuple[list, str]:
    """Dựng bookmark nhiều cấp từ trang mục lục in sẵn trong tài liệu."""
    ents, toc_pages = [], []
    for i in range(min(doc.page_count, scan_pages)):
        try:
            found = _toc_entries_on_page(doc[i])
        except Exception:
            found = []
        if found:
            toc_pages.append(i)
            ents += found
        elif toc_pages and i - toc_pages[-1] > 2:
            break
    if len(ents) < 4:
        return [], ""
    _levels_from_indent(ents, max_levels)
    ok = _resolve_toc_pages(doc, ents, toc_pages[-1] + 1)
    if ok < max(3, len(ents) * 0.5):
        return [], ""
    toc, nest = [], _Nester(max_levels)
    for e in ents:
        if e.get("page") is None:
            continue
        page = doc[e["page"]]
        y = 0.0
        try:
            hits = page.search_for(e["title"][:70])
            if hits:
                y = max(0.0, hits[0].y0 - 8)
        except Exception:
            pass
        lvl = nest.level(e["raw"])
        toc.append([lvl, e["title"], e["page"] + 1, {"kind": pymupdf.LINK_GOTO, "to": pymupdf.Point(0, y),
                                                     "collapse": lvl >= 2}])
    pages_txt = ", ".join(str(p + 1) for p in toc_pages[:6]) + ("…" if len(toc_pages) > 6 else "")
    return toc, f"trang mục lục in trong tài liệu (trang {pages_txt})"


def headings_from_fonts(doc, max_levels: int = 4, max_pages: int = 1500) -> tuple[list, str]:
    """Đoán tiêu đề nhiều cấp theo kiểu chữ: to hơn chữ nội dung, hoặc cùng cỡ nhưng in đậm đứng riêng một dòng."""
    rows_all = []       # (trang, hàng)
    weight: dict[float, int] = {}
    seen_text: dict[str, set] = {}
    n = min(doc.page_count, max_pages)
    for i in range(n):
        try:
            page = doc[i]
            if _toc_entries_on_page(page):
                continue          # bỏ trang mục lục in sẵn
            rows = _page_rows(page)
        except Exception:
            continue
        for r in rows:
            weight[r["size"]] = weight.get(r["size"], 0) + len(r["text"])
            seen_text.setdefault(_norm_text(r["text"]), set()).add(i)
            rows_all.append((i, r))
    if not weight:
        return [], ""
    body = max(weight.items(), key=lambda kv: kv[1])[0]
    repeat_limit = max(3, n * 0.2)
    cand = []
    for i, r in rows_all:
        t = r["text"]
        if not (3 <= len(t) <= 120) or not any(c.isalpha() for c in t):
            continue
        if len(seen_text.get(_norm_text(t), ())) >= repeat_limit:
            continue              # tiêu đề chạy đầu/cuối trang lặp lại trên nhiều trang
        bigger = r["size"] >= body * 1.08
        bold_line = r["bold"] and r["size"] >= body * 0.98 and len(t) <= 90 and not t.rstrip().endswith((".", ",", ";", ":"))
        if bigger or bold_line:
            cand.append((i, r))
    if not cand or len(cand) > 0.45 * len(rows_all) + 5:
        return [], ""
    styles = sorted({(r["size"], r["bold"]) for _i, r in cand}, key=lambda s: (-s[0], not s[1]))
    # gộp đậm/không đậm cùng cỡ khi quá nhiều kiểu, rồi lấy tối đa max_levels kiểu lớn nhất
    if len(styles) > max_levels:
        sizes = sorted({s[0] for s in styles}, reverse=True)[:max_levels]
        rank = lambda r: sizes.index(r["size"]) + 1 if r["size"] in sizes else None
    else:
        rank = lambda r: styles.index((r["size"], r["bold"])) + 1
    toc, last, nest = [], None, _Nester(max_levels)
    for i, r in cand:
        lvl = rank(r)
        if lvl is None:
            continue
        if last and last[0] == i and last[2] == (r["size"], r["bold"]) and 0 <= r["y"] - last[1] < r["size"] * 1.9:
            toc[-1][1] = (toc[-1][1] + " " + r["text"])[:150]      # tiêu đề xuống dòng: nối lại
            last = (i, r["y"], (r["size"], r["bold"]))
            continue
        lvl = nest.level(lvl)
        toc.append([lvl, r["text"], i + 1, {"kind": pymupdf.LINK_GOTO, "to": pymupdf.Point(0, max(0, r["y0"] - 8)),
                                            "collapse": lvl >= 2}])
        last = (i, r["y"], (r["size"], r["bold"]))
    return toc, "kiểu chữ của tiêu đề (cỡ chữ lớn hơn hoặc in đậm)"


def guess_headings(doc, max_levels: int = 4, source: str = "auto") -> tuple[list, str]:
    """Tự dựng bookmark nhiều cấp. source: 'auto' | 'toc' (trang mục lục in sẵn) | 'font' (kiểu chữ tiêu đề).
    Trả về (danh sách bookmark, mô tả nguồn đã dùng)."""
    max_levels = max(1, min(8, int(max_levels)))
    if source in ("auto", "toc"):
        toc, how = headings_from_printed_toc(doc, max_levels)
        if toc or source == "toc":
            return toc, how
    return headings_from_fonts(doc, max_levels)


def _flat_copy(src, cache: dict | None):
    """Bản sao đã gắn chết ghi chú/ô biểu mẫu vào nội dung trang, để co giãn trang mà không mất chúng."""
    key = id(src)
    if cache is not None and key in cache:
        return cache[key]
    flat = pymupdf.open(stream=src.tobytes(), filetype="pdf")
    try:
        flat.bake()
    except Exception:
        pass
    for pg in flat:  # trang đang xoay: đưa góc xoay vào nội dung để khi co giãn vẫn đúng chiều đang hiển thị
        if pg.rotation:
            try:
                pg.remove_rotation()
            except Exception:
                pass
    if cache is not None:
        cache[key] = flat
    return flat


def add_page_resized(dst, src, pno: int, size: str | None, cache: dict | None = None,
                     orientation: str = "portrait") -> bool:
    """Thêm trang pno của src vào dst. Nếu size là tên khổ giấy và trang khác khổ đó thì co giãn cho vừa
    (giữ tỉ lệ, canh giữa). orientation: 'portrait' / 'landscape' = mọi trang cùng một hướng;
    'auto' = hướng theo trang gốc. Trả về True nếu trang đã được co giãn."""
    r = src[pno].rect
    if size in PAPER_PT:
        pw, ph = PAPER_PT[size]
        land = r.width > r.height if orientation == "auto" else orientation == "landscape"
        if land:
            pw, ph = ph, pw
        if abs(r.width - pw) > 1.5 or abs(r.height - ph) > 1.5:
            flat = _flat_copy(src, cache)
            page = dst.new_page(width=pw, height=ph)
            k = min(pw / r.width, ph / r.height)
            tw, th = r.width * k, r.height * k
            x0, y0 = (pw - tw) / 2, (ph - th) / 2
            page.show_pdf_page(pymupdf.Rect(x0, y0, x0 + tw, y0 + th), flat, pno)
            try:  # giữ lại liên kết web (toạ độ theo trang đã co giãn)
                sp = src[pno]
                for ln in sp.get_links():
                    if ln.get("kind") == pymupdf.LINK_URI and ln.get("uri"):
                        fr = pymupdf.Rect(ln["from"]) * sp.rotation_matrix
                        fr.normalize()
                        page.insert_link({"kind": pymupdf.LINK_URI, "uri": ln["uri"],
                                          "from": pymupdf.Rect(x0 + fr.x0 * k, y0 + fr.y0 * k,
                                                               x0 + fr.x1 * k, y0 + fr.y1 * k)})
            except Exception:
                pass
            return True
    dst.insert_pdf(src, from_page=pno, to_page=pno)
    return False


def merge_docs(parts: list, progress=None, size: str | None = None, orientation: str = "portrait") -> bytes:
    """Gộp các phần [(doc, [chỉ số trang...]), ...] thành một PDF, trả về bytes.
    size: None = giữ nguyên khổ từng trang; 'A4', 'A3'... = đưa mọi trang về khổ đó."""
    dst = pymupdf.open()
    cache: dict = {}
    for n, (src, pages) in enumerate(parts):
        if size in PAPER_PT:
            for p in pages:
                add_page_resized(dst, src, p, size, cache, orientation)
        else:
            runs: list[list[int]] = []
            for p in pages:  # gom các trang liên tiếp để giữ liên kết, bookmark trong từng đoạn
                if runs and runs[-1][1] + 1 == p:
                    runs[-1][1] = p
                else:
                    runs.append([p, p])
            for a, b in runs:
                dst.insert_pdf(src, from_page=a, to_page=b)
        if progress:
            progress(n + 1, len(parts))
    data = dst.tobytes(garbage=3, deflate=True)
    dst.close()
    for f in cache.values():
        f.close()
    return data


# ------------------------------------------------------------ hình ảnh gốc trong trang
def _xobject_images(doc, xref: int, depth: int = 0, seen=None) -> list[tuple]:
    """Các ảnh nằm trong một Form XObject (kể cả lồng nhau): [(xref, rộng, cao)]."""
    seen = seen if seen is not None else set()
    if xref in seen or depth > 4:
        return []
    seen.add(xref)
    out = []
    try:
        kind, val = doc.xref_get_key(xref, "Resources/XObject")
        if kind == "xref":
            val = doc.xref_object(int(val.split()[0]))
    except Exception:
        return out
    for m in re.finditer(r"(\d+)\s+0\s+R", val or ""):
        x = int(m.group(1))
        try:
            sub = doc.xref_get_key(x, "Subtype")[1]
            if sub == "/Image":
                out.append((x, int(doc.xref_get_key(x, "Width")[1]), int(doc.xref_get_key(x, "Height")[1])))
            elif sub == "/Form":
                out += _xobject_images(doc, x, depth + 1, seen)
        except Exception:
            continue
    return out


def _annot_image_xref(page, info: dict) -> int:
    """Ảnh nằm trong ghi chú / con dấu / khung chữ ký (không thuộc nội dung trang): tìm đúng ảnh gốc của nó."""
    doc = page.parent
    r = info["bbox"]
    best = 0
    try:
        items = list(page.annots() or []) + list(page.widgets() or [])
    except Exception:
        items = []
    for an in items:
        try:
            if not pymupdf.Rect(an.rect).intersects(r):
                continue
            kind, val = doc.xref_get_key(an.xref, "AP/N")
            if kind != "xref":
                continue
            for x, w, h in _xobject_images(doc, int(val.split()[0])):
                if (w, h) == (info.get("width"), info.get("height")):
                    return x
                best = best or x
        except Exception:
            continue
    return best


def page_images(page) -> list[dict]:
    """Các hình ảnh nằm trong nội dung trang: [{'bbox': Rect, 'xref': int, 'width': px, 'height': px}, ...]."""
    out = []
    try:
        infos = page.get_image_info(xrefs=True)
    except Exception:
        return out
    for d in infos:
        r = pymupdf.Rect(d["bbox"])
        if r.is_empty or r.width < 4 or r.height < 4:
            continue
        info = {"bbox": r, "xref": d.get("xref", 0), "width": d.get("width", 0), "height": d.get("height", 0)}
        if not info["xref"]:          # ảnh chèn dạng con dấu / chữ ký: nằm trong ghi chú chứ không trong nội dung trang
            info["xref"] = _annot_image_xref(page, info)
            info["annot"] = True
        out.append(info)
    return out


def image_at(images: list[dict], pt) -> dict | None:
    """Ảnh tại điểm pt; nếu chồng nhau thì lấy ảnh nhỏ nhất (thường là ảnh nằm trên)."""
    pt = pymupdf.Point(pt)
    hits = [d for d in images if pt in d["bbox"]]
    return min(hits, key=lambda d: d["bbox"].get_area()) if hits else None


def image_pixmap(doc, page, info: dict):
    """Lấy ảnh gốc (đúng độ phân giải trong file) dạng Pixmap RGB hoặc RGBA."""
    xref = info.get("xref", 0)
    pix = None
    if xref:
        try:
            pix = pymupdf.Pixmap(doc, xref)
            smask = doc.extract_image(xref).get("smask", 0)
            if smask:
                mask = pymupdf.Pixmap(doc, smask)
                if pix.alpha:
                    pix = pymupdf.Pixmap(pix, 0)
                pix = pymupdf.Pixmap(pix, mask)
            if pix.colorspace is None or pix.colorspace.n != 3:
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        except Exception:
            pix = None
    if pix is None:  # ảnh nhúng trực tiếp hoặc định dạng lạ: chụp lại đúng vùng ảnh trên trang
        r = info["bbox"]
        z = max(1.0, min(6.0, (info.get("width") or 0) / max(1.0, r.width))) if info.get("width") else 3.0
        clip = r * page.rotation_matrix if page.rotation else r
        clip.normalize()
        pix = page.get_pixmap(matrix=pymupdf.Matrix(z, z), clip=clip, alpha=False, annots=bool(info.get("annot")))
    return pix


def save_image(doc, page, info: dict, path: str) -> None:
    ext = os.path.splitext(path)[1].lower()
    xref = info.get("xref", 0)
    if xref and ext in (".jpg", ".jpeg"):
        try:
            raw = doc.extract_image(xref)
            if raw.get("ext") in ("jpeg", "jpg") and not raw.get("smask") and raw.get("colorspace") == 3:
                Path(path).write_bytes(raw["image"])   # ghi thẳng dữ liệu gốc, không nén lại
                return
        except Exception:
            pass
    pix = image_pixmap(doc, page, info)
    if ext in (".jpg", ".jpeg"):
        if pix.alpha:
            pix = pymupdf.Pixmap(pix, 0)
        pix.save(path, jpg_quality=95)
    else:
        pix.save(path)


def delete_image(page, info: dict) -> None:
    """Gỡ một hình ảnh khỏi trang (chữ và nét vẽ xung quanh giữ nguyên)."""
    r = info["bbox"]
    if info.get("annot"):        # ảnh chèn dạng con dấu / chữ ký: xoá chính ghi chú đó (khung chữ ký số thì không đụng tới)
        for an in list(page.annots() or []):
            if an.type[1] == "Stamp" and (pymupdf.Rect(an.rect) & r).get_area() > 0.5 * r.get_area():
                page.delete_annot(an)
                return
        raise ValueError("Hình này thuộc khung chữ ký số hoặc ô biểu mẫu nên không xoá riêng được.")
    c = (r.tl + r.br) / 2
    w, h = max(1.0, r.width * 0.1), max(1.0, r.height * 0.1)
    page.add_redact_annot(pymupdf.Rect(c.x - w, c.y - h, c.x + w, c.y + h), fill=False, cross_out=False)
    page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                          text=pymupdf.PDF_REDACT_TEXT_NONE)


def export_all_images(doc, out_dir: str, base: str, min_size: int = 32, progress=None) -> int:
    """Xuất mọi hình ảnh gốc trong tài liệu ra thư mục. Trả về số ảnh đã ghi."""
    seen, n = set(), 0
    for i in range(doc.page_count):
        page = doc[i]
        for k, info in enumerate(page_images(page), 1):
            x = info["xref"]
            if x and x in seen:
                continue
            if x:
                seen.add(x)
            if min(info["width"] or 0, info["height"] or 0) < min_size:
                continue
            ext = ".png"
            try:
                raw = doc.extract_image(x) if x else {}
                if raw.get("ext") in ("jpeg", "jpg") and not raw.get("smask") and raw.get("colorspace") == 3:
                    ext = ".jpg"
            except Exception:
                pass
            try:
                save_image(doc, page, info, os.path.join(out_dir, f"{base}_trang{i + 1:03d}_anh{k:02d}{ext}"))
                n += 1
            except Exception:
                continue
        if progress:
            progress(i + 1, doc.page_count)
    return n


# ----------------------------------------------------------------- tìm chữ
def _fold(s: str, plain: bool) -> str:
    """Chuẩn hoá để so khớp: không phân biệt hoa thường (cả chữ có dấu: Đ/đ, Ạ/ạ…); plain = bỏ luôn dấu."""
    s = unicodedata.normalize("NFC", s).casefold()
    if plain:
        s = "".join(ch for ch in unicodedata.normalize("NFD", s) if not unicodedata.combining(ch)).replace("đ", "d")
    return s


def _is_cjk(c: str) -> bool:
    o = ord(c)
    return (0x2E80 <= o <= 0x9FFF or 0xAC00 <= o <= 0xD7AF or 0x1100 <= o <= 0x11FF or 0xF900 <= o <= 0xFAFF
            or 0xFF00 <= o <= 0xFFEF or 0x20000 <= o <= 0x2FA1F)


def _drop_cjk_spaces(chars: list, owners: list | None = None):
    """Bỏ khoảng trắng / chỗ ngắt dòng nằm giữa hai chữ Trung–Nhật–Hàn. Chữ OCR hay bị chèn dấu cách giả
    hoặc ngắt dòng giữa từ ("增值 税"), còn chữ Hàn thì lúc có lúc mất dấu cách: bỏ đi ở cả trang lẫn từ cần tìm."""
    keep = [i for i, c in enumerate(chars)
            if not (c == " " and 0 < i < len(chars) - 1 and _is_cjk(chars[i - 1]) and _is_cjk(chars[i + 1]))]
    return [chars[i] for i in keep], ([owners[i] for i in keep] if owners is not None else None)


def ocr_page_words(page, lang: str, dpi: int = 300) -> list:
    """Nhận dạng chữ của trang bằng Tesseract (không sửa trang). Trả về danh sách từ giống page.get_text("words"),
    toạ độ theo trang gốc."""
    td = find_tessdata()
    if not td:
        return []
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    one = pymupdf.open("pdf", pix.pdfocr_tobytes(language=lang, tessdata=td))
    try:
        op = one[0]
        # trang OCR có cỡ đúng bằng trang đang hiển thị (đã xoay nếu trang có xoay): đổi về toạ độ trang gốc
        kx = page.rect.width / max(1e-6, op.rect.width)
        ky = page.rect.height / max(1e-6, op.rect.height)
        m = page.derotation_matrix if page.rotation else None
        out = []
        for w in op.get_text("words"):
            r = pymupdf.Rect(w[0] * kx, w[1] * ky, w[2] * kx, w[3] * ky)
            if m is not None:
                r = (r * m).normalize()
            out.append((r.x0, r.y0, r.x1, r.y1) + tuple(w[4:]))
        return out
    finally:
        one.close()


def search_words(words: list, query: str) -> list:
    """Như search_page nhưng tìm trong danh sách từ (vd. chữ tự nhận dạng bằng OCR)."""
    plain = _fold(query, True) == _fold(query, False)
    q = " ".join(_fold(query, plain).split())
    if not q:
        return []
    cjk = any(_is_cjk(c) for c in q)
    if cjk:
        q = "".join(_drop_cjk_spaces(list(q))[0])
    flat, own = [], []
    for k, w in enumerate(words):
        if flat:
            flat.append(" ")
            own.append(-1)
        for c in _fold(w[4], plain):
            flat.append(c)
            own.append(k)
    if cjk:
        flat, own = _drop_cjk_spaces(flat, own)
    hay = "".join(flat)
    out, start = [], 0
    while True:
        k = hay.find(q, start)
        if k < 0:
            break
        start = k + 1
        groups: dict = {}
        for o in own[k:k + len(q)]:
            if o < 0:
                continue
            w = words[o]
            key = (w[5], w[6]) if len(w) > 6 else 0
            r = pymupdf.Rect(w[:4])
            groups[key] = groups[key] | r if key in groups else r
        out += list(groups.values())
    return out


def search_page(page, query: str) -> list:
    """Tìm `query` trong trang, trả về các khung chữ nhật (mỗi dòng của một kết quả một khung).
    Khác search_for của MuPDF (chỉ không phân biệt hoa thường với chữ không dấu): "đại việt" tìm ra "ĐẠI VIỆT".
    Gõ không dấu ("dai viet") thì tìm cả chữ có dấu."""
    plain = _fold(query, True) == _fold(query, False)
    q = " ".join(_fold(query, plain).split())
    if not q:
        return []
    cjk = any(_is_cjk(c) for c in q)
    if cjk:
        q = "".join(_drop_cjk_spaces(list(q))[0])
    text: list[str] = []      # từng mảnh chữ đã chuẩn hoá
    owner: list[int] = []     # mảnh đó thuộc ký tự nào (-1: khoảng trắng chèn giữa các dòng)
    chars: list[tuple] = []   # (khung, số dòng)
    line_no = 0
    for b in _fix_spans(page, page.get_text("rawdict", flags=pymupdf.TEXT_USE_GID_FOR_UNKNOWN_UNICODE)).get("blocks", []):
        for ln in b.get("lines", []):
            line_no += 1
            text.append(" ")
            owner.append(-1)
            for sp in ln.get("spans", []):
                for ch in sp.get("chars", []):
                    f = _fold(ch["c"], plain)
                    if not f or f.isspace():
                        f = " "
                    chars.append((ch["bbox"], line_no))
                    for piece in f:
                        text.append(piece)
                        owner.append(len(chars) - 1)
    # gộp khoảng trắng liên tiếp thành một
    flat, own = [], []
    for c, o in zip(text, owner):
        if c == " " and flat and flat[-1] == " ":
            continue
        flat.append(c)
        own.append(o)
    if cjk:
        flat, own = _drop_cjk_spaces(flat, own)
    hay = "".join(flat)
    out, start = [], 0
    while True:
        k = hay.find(q, start)
        if k < 0:
            break
        start = k + 1
        groups: dict[int, pymupdf.Rect] = {}
        for o in own[k:k + len(q)]:
            if o < 0:
                continue
            bb, ln = chars[o]
            r = pymupdf.Rect(bb)
            if r.is_empty:
                continue
            groups[ln] = groups[ln] | r if ln in groups else r
        out += [groups[g] for g in sorted(groups)]
    return out


# ----------------------------------------------------------------- dời / co giãn / xoay hình ảnh gốc
class ImageEditError(Exception):
    pass


def _fmt(v: float) -> str:
    s = f"{v:.5f}".rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s


def transform_image(page, info: dict, a_m) -> None:
    """Đặt lại một hình ảnh gốc của trang theo phép biến đổi `a_m` (toạ độ trang: dời, co giãn, xoay quanh một điểm).
    Chỉ sửa đúng lệnh vẽ ảnh đó trong nội dung trang (giữ nguyên thứ tự lớp: chữ đè lên ảnh vẫn đè lên), ảnh gốc giữ
    nguyên chất lượng, không phải nhúng lại."""
    doc = page.parent
    xref = info.get("xref", 0)
    if not xref or info.get("annot"):
        raise ImageEditError("Không tìm thấy ảnh trong nội dung trang.")
    infos = [d for d in page.get_image_info(xrefs=True) if d.get("xref") == xref]
    target = pymupdf.Rect(info["bbox"])
    k = min(range(len(infos)), key=lambda j: abs(pymupdf.Rect(infos[j]["bbox"]).x0 - target.x0)
            + abs(pymupdf.Rect(infos[j]["bbox"]).y0 - target.y0)
            + abs(pymupdf.Rect(infos[j]["bbox"]).x1 - target.x1)
            + abs(pymupdf.Rect(infos[j]["bbox"]).y1 - target.y1), default=None)
    if k is None:
        raise ImageEditError("Không tìm thấy ảnh trong nội dung trang.")
    by_ref: dict[int, set] = {}      # nơi khai báo ảnh: 0 = nội dung trang, số khác = nhóm nội dung (form) chứa ảnh
    for it in page.get_images(full=True):
        if it[0] == xref:
            by_ref.setdefault(it[9], set()).add(it[7])

    def find(data: bytes, names: set):
        pat = re.compile(rb"/(" + b"|".join(re.escape(n.encode("latin-1")) for n in sorted(names, key=len, reverse=True))
                         + rb")\s+Do\b")
        return [(m.start(), m.end(), m.group(1)) for m in pat.finditer(data)]

    hits = []        # (xref luồng, vị trí đầu, vị trí cuối, tên)
    for sx in page.get_contents():
        if 0 in by_ref:
            hits += [(sx,) + h for h in find(doc.xref_stream(sx) or b"", by_ref[0])]
    form_hits = []
    for ref, names in by_ref.items():
        if ref:
            form_hits += [(ref,) + h for h in find(doc.xref_stream(ref) or b"", names)]
    if form_hits:
        # ảnh nằm trong nhóm nội dung (file xuất từ Word, trình duyệt…): sửa được khi trên trang chỉ có đúng một chỗ vẽ ảnh đó
        if len(infos) == 1 and len(hits) + len(form_hits) == 1:
            hits = form_hits
        else:
            raise ImageEditError("Ảnh này được vẽ nhiều lần qua một nhóm nội dung dùng chung nên chưa sửa riêng được.")
    if len(hits) != len(infos):
        raise ImageEditError("Ảnh này được vẽ qua một nhóm nội dung phức tạp nên chưa sửa riêng được.")
    T = pymupdf.Matrix(infos[k]["transform"])
    F = pymupdf.Matrix(1, 0, 0, -1, 0, 1)            # ảnh trong MuPDF tính từ góc trên, trong PDF tính từ góc dưới
    try:
        X = F * T * pymupdf.Matrix(a_m) * ~T * F
    except Exception:
        raise ImageEditError("Ảnh quá mỏng để biến đổi.")
    sx, a, b, name = hits[k]
    data = doc.xref_stream(sx)
    rep = b"q " + " ".join(_fmt(v) for v in (X.a, X.b, X.c, X.d, X.e, X.f)).encode() + b" cm /" + name + b" Do Q"
    doc.update_stream(sx, data[:a] + rep + data[b:])


def image_quad(info: dict, a_m) -> pymupdf.Quad:
    """Bốn góc của ảnh sau phép biến đổi (để vẽ khung xem trước)."""
    return pymupdf.Rect(info["bbox"]).quad * pymupdf.Matrix(a_m)
